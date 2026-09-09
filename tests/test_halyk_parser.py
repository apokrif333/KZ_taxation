from __future__ import annotations

import tempfile
import unittest
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook

from conftest_imports import SRC  # noqa: F401
from kztax270.brokers.account_detection import detect_report_account_id, detect_report_period_end
from kztax270.brokers.halyk import HalykParser, parse_halyk_xlsx
from kztax270.brokers.registry import default_registry
from kztax270.canonical.validation import validate_dataset_for_tax_forms
from kztax270.reconciliation.engine import ReconciliationEngine
from kztax270.reconciliation.models import ReconciliationSeverity
from kztax270.reference.fx import AnnualFxRateProvider

MOVEMENT_HEADERS = (
    "№",
    "Наименование (тип)",
    "Рынок ЦБ",
    "Тип ЦБ",
    "ISIN ЦБ",
    "Направление операции",
    "Перевод откуда",
    "Перевод куда",
    "Код страны регистрации эмитента",
    "Дата торговой операции",
    "Количество ЦБ",
    "Код валюты",
    "ISIN конвертируемой ЦБ",
    "Цена приобретения одной ЦБ",
    "Коэффициент конвертации",
    "Комиссии по сделке",
    "Номер заказа",
)


def _report(path: Path, account_id: str = "102TEST") -> None:
    workbook = Workbook()
    movement = workbook.active
    movement.title = "Движение ФИ"
    movement.cell(1, 1, "Отчет по движению финансовых инструментов")
    movement.cell(4, 1, "Номер лицевого счета")
    movement.cell(4, 6, account_id)
    movement.cell(5, 1, "По состоянию на")
    movement.cell(5, 6, "31.12.2025")
    for column, value in enumerate(MOVEMENT_HEADERS, start=1):
        movement.cell(8, column, value)
    movement.append((1, "Остатки денежных средств", None, None, None, None, None, None, None, "01.01.1900", 0, "USD", None, 100, 0, 0, 0))
    movement.append((2, "TBILL, Дисконтные облигации", None, "Дисконтные облигации государственные иностранные", "US0000000001", "Прием инструментов при покупке", "Dealer", "-", "US", "01.02.2025", 10, "USD", None, 95, 0, 520, 101))
    movement.append((2, "TBILL, Дисконтные облигации", "NON EXCHANGE", "Дисконтные облигации государственные иностранные", "US0000000001", "Списание инструментов при погашении ЦБ", "-", "NON EXCHANGE (58)", "US", "01.12.2025", -10, "USD", None, 100, 0, 0, 102))
    movement.append((3, "HSBK, Простые акции", "KASE", "Акции простые резидентов РК", "KZ0000000001", "Прием инструментов при покупке", "KASE (52)", "-", "KZ", "15.12.2025", 2, "KZT", None, 200, 0, 10, 103))

    income = workbook.create_sheet("Дивиденды")
    income.cell(4, 1, "Номер лицевого счета")
    income.cell(4, 6, account_id)
    income.cell(5, 1, "Период за который предоставляется отчет с/ по :")
    income.cell(5, 6, "01.01.2025 - 31.12.2025")
    headers = (
        "№",
        "Дата транзакции",
        "Наименование (тип)",
        "Количество ЦБ",
        "Код валюты",
        "Сумма в валюте",
        "Тип операции",
        "Тип ЦБ",
        "ISIN ЦБ",
        "Рынок ЦБ",
        "Код страны регистрации эмитента",
    )
    for column, value in enumerate(headers, start=1):
        income.cell(8, column, value)
    income.append((1, "15.06.2025", "TBILL, Дисконтные облигации", 10, "USD", 10, "Выплата купона", "Дисконтные облигации государственные иностранные", "US0000000001", "OTC", "US"))

    repo = workbook.create_sheet("РЕПО")
    repo.cell(4, 1, "Номер лицевого счета")
    repo.cell(4, 6, account_id)
    repo.cell(8, 1, "№")
    repo.cell(8, 2, "Наименование (тип)")
    repo.cell(8, 3, "Дата транзакции")
    repo.cell(8, 4, "Количество ЦБ")
    repo.cell(8, 5, "Код валюты")
    repo.cell(8, 6, "Объем открытия в валюте")
    repo.cell(8, 7, "Объем закрытия в валюте")

    positions = workbook.create_sheet("ЦБ")
    positions.cell(4, 1, "Номер лицевого счета")
    positions.cell(4, 5, account_id)
    positions.cell(5, 1, "По состоянию на")
    positions.cell(5, 5, "31.12.2025")
    position_headers = (
        "№",
        "Наименование",
        "ISIN ЦБ",
        "Тип ЦБ",
        "Код страны регистрации эмитента",
        "Количество ЦБ",
        "Адрес регистрации",
        "Код валюты",
    )
    for column, value in enumerate(position_headers, start=1):
        positions.cell(8, column, value)
    positions.append((1, "HSBK, Простые акции", "KZ0000000001", "Акции простые резидентов РК", "KZ", 2, "Казахстан", "KZT"))
    workbook.save(path)
    workbook.close()


class HalykParserTests(unittest.TestCase):
    def test_xlsx_sections_and_metadata_are_parsed(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "102TEST tax report.xlsx"
            _report(path)
            report = parse_halyk_xlsx(path)

        self.assertEqual(report.account_id, "102TEST")
        self.assertEqual(report.period_end.isoformat(), "2025-12-31")
        self.assertEqual(len(report.trades), 3)
        self.assertEqual(report.trades[1]["operation_kind"], "redemption")
        self.assertEqual(report.cash_balances[0]["ending_cash"], "100")
        self.assertEqual(report.positions[0]["quantity"], "2")
        self.assertEqual(report.income[0]["amount"], "10")

    def test_dataset_builds_fifo_income_cash_and_reconciliation(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            raw_root = Path(tmp)
            broker_root = raw_root / "halyk"
            broker_root.mkdir()
            path = broker_root / "102TEST tax report.xlsx"
            _report(path)
            parser = HalykParser(AnnualFxRateProvider({(2025, "USD"): Decimal("520")}))
            reports = parser.discover_reports(raw_root, "102TEST")
            result = parser.parse_reports(reports, "102TEST")

        dataset = result.dataset
        validate_dataset_for_tax_forms(dataset)
        self.assertEqual(dataset.metadata.broker, "halyk")
        self.assertEqual(len(dataset.tables["Trades"]), 3)
        purchase = dataset.tables["Trades"][0]
        self.assertEqual(Decimal(purchase["commission"]), Decimal("1"))
        self.assertEqual(Decimal(purchase["amount_with_commission"]), Decimal("951"))
        self.assertEqual(dataset.tables["Fifo"][0]["operation_type"], "bond_redemption")
        self.assertEqual(dataset.tables["CorporateActions"][0]["action_type"], "maturity")
        self.assertEqual(dataset.tables["Coupons"][0]["gross_amount"], "10.00")
        self.assertEqual(dataset.tables["CashBalances"][0]["ending_cash"], "100.00")
        self.assertEqual(dataset.tables["Unprocessed"], [])
        self.assertTrue(
            any(row["table"] == "Yearly Bonds Redemption" for row in dataset.tables["Years_Results"])
        )
        reconciliation = ReconciliationEngine().reconcile_dataset(dataset)
        self.assertEqual(
            [row for row in reconciliation if row.severity == ReconciliationSeverity.ERROR],
            [],
        )

    def test_same_isin_custody_credit_is_costed_transfer_in(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            raw_root = Path(tmp)
            broker_root = raw_root / "halyk"
            broker_root.mkdir()
            path = broker_root / "102TEST tax report.xlsx"
            _report(path)
            workbook = load_workbook(path)
            movement = workbook["Движение ФИ"]
            movement.append(
                (
                    4,
                    "HSBK, Простые акции",
                    None,
                    "Акции простые резидентов РК",
                    "KZ0000000001",
                    "Приказ на зачисление ЦБ клиент ХФ <- Контрагент ЦБ б/сделки",
                    "Контрагент (383)",
                    "-",
                    "KZ",
                    "14.12.2025",
                    20,
                    None,
                    "KZ0000000001",
                    0,
                    154.5,
                    0,
                    104,
                )
            )
            workbook["ЦБ"].cell(9, 6, 22)
            workbook.save(path)
            workbook.close()

            parser = HalykParser(AnnualFxRateProvider({(2025, "USD"): Decimal("520")}))
            result = parser.parse_reports(parser.discover_reports(raw_root, "102TEST"), "102TEST")

        dataset = result.dataset
        self.assertEqual(len(dataset.tables["Trades"]), 3)
        self.assertEqual(len(dataset.tables["Transfers"]), 1)
        transfer = dataset.tables["Transfers"][0]
        self.assertEqual(transfer["direction"], "in")
        self.assertEqual(transfer["isin"], "KZ0000000001")
        self.assertEqual(Decimal(transfer["quantity"]), Decimal("20"))
        self.assertEqual(Decimal(transfer["price"]), Decimal("154.5"))
        hsbk_quantity = sum(
            Decimal(row["quantity"]) for row in dataset.tables["Positions"] if row["isin"] == "KZ0000000001"
        )
        self.assertEqual(hsbk_quantity, Decimal("22"))
        self.assertEqual(dataset.tables["Unprocessed"], [])
        reconciliation = ReconciliationEngine().reconcile_dataset(dataset)
        self.assertEqual(
            [row for row in reconciliation if row.severity == ReconciliationSeverity.ERROR],
            [],
        )

    def test_security_currency_change_closes_one_fifo_position(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            raw_root = Path(tmp)
            broker_root = raw_root / "halyk"
            broker_root.mkdir()
            path = broker_root / "102TEST tax report.xlsx"
            _report(path)
            workbook = load_workbook(path)
            movement = workbook["Движение ФИ"]
            movement.append(
                (4, "US_INTC, Common Stocks", "KASE", "Акции простые иностранные", "US4581401001", "Прием инструментов при покупке", "KASE (52)", "-", "US", "29.09.2021", 1, "KZT", None, 23700, 0, 0, 104)
            )
            movement.append(
                (5, "US_INTC, Common Stocks", "KASE", "Акции простые иностранные", "US4581401001", "Прием инструментов при покупке", "KASE (52)", "-", "US", "14.10.2021", 1, "KZT", None, 23993, 0, 0, 105)
            )
            movement.append(
                (6, "INTC_KZ, Common Stocks", "KASE", "Акции простые иностранные", "US4581401001", "Прием инструментов при покупке", "KASE (52)", "-", "US", "15.08.2022", 5, "USD", None, 35.79, 0, 0, 106)
            )
            movement.append(
                (7, "INTC_KZ, Common Stocks", "KASE", "Акции простые иностранные", "US4581401001", "Перечисление инструментов при продаже", "-", "KASE (52)", "US", "25.01.2024", -7, "USD", None, 48.25, 0, 0, 107)
            )
            workbook.save(path)
            workbook.close()

            parser = HalykParser(
                AnnualFxRateProvider(
                    {
                        (2021, "USD"): Decimal("400"),
                        (2022, "USD"): Decimal("450"),
                        (2024, "USD"): Decimal("500"),
                        (2025, "USD"): Decimal("520"),
                    }
                )
            )
            result = parser.parse_reports(parser.discover_reports(raw_root, "102TEST"), "102TEST")

        dataset = result.dataset
        intc_positions = [
            row
            for row in dataset.tables["Positions"]
            if row["isin"] == "US4581401001" and row["year"] == 2025
        ]
        intc_fifo = [row for row in dataset.tables["Fifo"] if row["isin"] == "US4581401001"]
        self.assertEqual(intc_positions, [])
        self.assertEqual(Decimal(sum(Decimal(row["exit_quantity"]) for row in intc_fifo)), Decimal("7"))
        self.assertTrue(all(row["currency"] == "USD" for row in intc_fifo))
        reconciliation = ReconciliationEngine().reconcile_dataset(dataset)
        self.assertEqual(
            [row for row in reconciliation if row.severity == ReconciliationSeverity.ERROR],
            [],
        )

    def test_account_detection_and_registry_expose_halyk(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "report.xlsx"
            _report(path, "1028100VKP")
            self.assertEqual(detect_report_account_id("halyk", path), "1028100VKP")
            self.assertEqual(detect_report_period_end("halyk", path).isoformat(), "2025-12-31")
        self.assertIn("halyk", default_registry().broker_codes())


if __name__ == "__main__":
    unittest.main()
