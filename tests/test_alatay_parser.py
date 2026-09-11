from __future__ import annotations

import tempfile
import unittest
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook

from conftest_imports import SRC  # noqa: F401
from kztax270.brokers.account_detection import detect_report_account_id
from kztax270.brokers.alatay import (
    AlatayParser,
    ParsedAlatayReport,
    _parse_trade_row,
    build_canonical_dataset,
    parse_alatay_csv,
    parse_alatay_xlsx,
)
from kztax270.brokers.registry import default_registry
from kztax270.canonical.validation import validate_dataset_for_tax_forms
from kztax270.reconciliation.engine import ReconciliationEngine
from kztax270.reconciliation.models import ReconciliationMetric, ReconciliationSeverity
from kztax270.reference.fx import AnnualFxRateProvider

REPORT_2023 = """Отчет движения ценных бумаг
,
,ФИО/Наименование клиента:,ТЕСТОВЫЙ КЛИЕНТ
,№ лицевого счета:,ATEST
,Отчет составлен на:,01.01.2023-31.12.2023
,
Ценные бумаги в портфеле клиента на конец отчетного периода:
№,Эмитент,Вид ЦБ,ISIN,Номинал,Количество ЦБ,Код страны регистрации эмитента
1,ЭМИТЕНТ А,АКЦИИ,KZ0000000001,1,10,
,
Движение Ценных бумаг клиента за отчетный период:
Дата расчетов сделки/операции,Эмитент,Вид ЦБ,ISIN,Номинал,Вид сделки/тип операции,Количество ЦБ,Цена за 1 ЦБ (% от номинала облигации),Валюта сделки,Сумма сделки/операции,Рынок заключения сделки,Код страны регистрации эмитента,Сумма комиссии
01.02.2023,ЭМИТЕНТ А,АКЦИИ,KZ0000000001,1,Покупка,10,100,KZT,1 000,KASE_MOEX,,2
"""


REPORT_2024 = """Отчет движения ценных бумаг
,
,ФИО/Наименование клиента:,ТЕСТОВЫЙ КЛИЕНТ
,№ лицевого счета:,ATEST
,Отчет составлен на:,01.01.2024-31.12.2024
,
Ценные бумаги в портфеле клиента на конец отчетного периода:
№,Эмитент,Вид ЦБ,ISIN,Номинал,Количество ЦБ,Код страны регистрации эмитента
1,ЭМИТЕНТ А,АКЦИИ,KZ0000000001,1,7,
2,ЭМИТЕНТ Б,ОБЛИГАЦИИ,KZ2000000002,1 000,5,
,
Движение Ценных бумаг клиента за отчетный период:
Дата расчетов сделки/операции,Эмитент,Вид ЦБ,ISIN,Номинал,Вид сделки/тип операции,Количество ЦБ,Цена за 1 ЦБ (% от номинала облигации),Валюта сделки,Сумма сделки/операции,Рынок заключения сделки,Код страны регистрации эмитента,Сумма комиссии
01.03.2024,ЭМИТЕНТ А,АКЦИИ,KZ0000000001,1,Продажа,3,120,KZT,360,KASE_MOEX,,1
02.03.2024,ЭМИТЕНТ Б,ОБЛИГАЦИИ,KZ2000000002,1 000,Покупка,5,1 050,KZT,5 250,KASE_MOEX,,0
"""


HISTORICAL_SECURITIES_REPORT = """Ценные бумаги на начало отчетного периода и их движение до начала отчетного периода
,
,ФИО/Наименование клиента:,ТЕСТОВЫЙ КЛИЕНТ
,№ лицевого счета:,00123
,Отчет составлен на:,01.01.2025
,
Ценные бумаги в портфеле клиента на начало отчетного периода:
№,Эмитент,Вид ЦБ,ISIN,Номинал,Количество ЦБ
1,ЭМИТЕНТ А,АКЦИИ,KZ0000000001,1,10
,
Сделки по ценным бумагам, указанным на начало отчетного периода, за время до начала отчетного периода:
Дата расчетов сделки/операции,Эмитент,Вид ЦБ,ISIN,Номинал,Вид сделки/тип операции,Количество ЦБ,Цена за 1 ЦБ (% от номинала облигации),Валюта сделки,Сумма сделки/операции,Рынок заключения сделки,Сумма комиссии
01.02.2023,ЭМИТЕНТ А,АКЦИИ,KZ0000000001,1,Покупка /Народное IPO,10,100,KZT,1 000,KASE_MOEX,2
02.02.2023,ЭМИТЕНТ А,АКЦИИ,KZ0000000001,1,Перевод основной (получатель без смены прав собственности),10,1,KZT,10,ЦД ЦБ,0
02.02.2023,ЭМИТЕНТ А,АКЦИИ,KZ0000000001,1,Перевод основной (поставщик без смены прав собственности),10,1,KZT,10,ЦД ЦБ,0
"""


CASH_REPORT = """ОТЧЕТ ДВИЖЕНИЯ ДЕНЕЖНЫХ СРЕДСТВ
,
,ФИО/Наименование клиента:,ТЕСТОВЫЙ КЛИЕНТ
,№ лицевого счета:,123
,Отчет составлен на:,01.01.2024-31.12.2024
,
Входящий остаток на начало периода:,100 KZT
Исходящий остаток на конец периода:,141 KZT
,
Дата проведения операции/сделки,Содержание операции/сделки,Тип операции/вид сделки,Эмитент,ISIN,Входящий остаток,Приход,Расход,Исходящий остаток,Код валюты,Тип ЦБ,Код страны регистрации эмитента,Наименование рынка
01.06.2024,Зачисление денежных средств,Зачисление вознаграждения (дивиденды),ЭМИТЕНТ А,KZ0000000001,100,10,0,110,KZT,АКЦИИ,KZ,KASE
01.07.2024,Зачисление денежных средств,Зачисление вознаграждения (купон),ЭМИТЕНТ Б,KZ2000000002,110,1,0,111,KZT,ОБЛИГАЦИИ,KZ,KASE
01.08.2024,Зачисление денежных средств по сделке,Погашение ЦБ KZ2000000002,ЭМИТЕНТ Б,KZ2000000002,111,30,0,141,KZT,ОБЛИГАЦИИ,KZ,KASE
"""


class AlatayParserTests(unittest.TestCase):
    def test_csv_metadata_positions_and_trades_are_parsed(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "ATEST alatay 2023.csv"
            path.write_text(REPORT_2023, encoding="utf-8")
            report = parse_alatay_csv(path)

        self.assertEqual(report.account_id, "ATEST")
        self.assertEqual(report.period_end.isoformat(), "2023-12-31")
        self.assertEqual(report.positions[0]["quantity"], "10")
        self.assertEqual(report.trades[0]["price"], "100")
        self.assertEqual(report.trades[0]["amount"], "1000")
        self.assertEqual(report.trades[0]["commission"], "2")

    def test_history_builds_fifo_positions_and_yearly_reconciliation(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            raw_root = Path(tmp)
            broker_root = raw_root / "alatay"
            broker_root.mkdir()
            (broker_root / "ATEST alatay 2023.csv").write_text(REPORT_2023, encoding="utf-8")
            (broker_root / "ATEST alatay 2024.csv").write_text(REPORT_2024, encoding="utf-8")
            (broker_root / "OTHER alatay 2024.csv").write_text(REPORT_2024, encoding="utf-8")

            parser = AlatayParser()
            reports = parser.discover_reports(raw_root, "ATEST")
            result = parser.parse_reports(reports, "ATEST")

        self.assertEqual([report.path.name for report in reports], ["ATEST alatay 2023.csv", "ATEST alatay 2024.csv"])
        validate_dataset_for_tax_forms(result.dataset)
        self.assertEqual([row["quantity"] for row in result.dataset.tables["Trades"]], ["10", "-3", "5"])
        self.assertEqual(len(result.dataset.tables["Fifo"]), 1)
        self.assertEqual(result.dataset.tables["Fifo"][0]["enter_price"], "100")
        self.assertEqual(result.dataset.tables["Fifo"][0]["exit_price"], "120")
        self.assertEqual(result.dataset.tables["Unprocessed"], [])

        reconciliation = ReconciliationEngine().reconcile_dataset(result.dataset)
        position_rows = [
            row for row in reconciliation if row.metric == ReconciliationMetric.ENDING_POSITION_QUANTITY
        ]
        self.assertEqual(len(position_rows), 3)
        self.assertTrue(all(row.severity == ReconciliationSeverity.INFO for row in position_rows))
        by_key = {(row.year, row.instrument_key): row for row in position_rows}
        self.assertEqual(by_key[(2023, "KZ0000000001")].canonical_value, Decimal("10"))
        self.assertEqual(by_key[(2024, "KZ0000000001")].canonical_value, Decimal("7"))
        self.assertEqual(by_key[(2024, "KZ2000000002")].canonical_value, Decimal("5"))
        self.assertEqual(
            [row for row in reconciliation if row.severity == ReconciliationSeverity.ERROR],
            [],
        )

        instruments = {row["isin"]: row for row in result.dataset.tables["Instruments"]}
        self.assertEqual(instruments["KZ0000000001"]["listing_exchange"], "KASE")
        self.assertEqual(instruments["KZ2000000002"]["type"], "Bonds")

    def test_historical_securities_and_cash_reports_are_combined(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            raw_root = Path(tmp)
            broker_root = raw_root / "alatay"
            broker_root.mkdir()
            (broker_root / "00123 ОДЦБ.csv").write_text(HISTORICAL_SECURITIES_REPORT, encoding="utf-8")
            (broker_root / "00123 ОДДС.csv").write_text(CASH_REPORT, encoding="utf-8")

            parser = AlatayParser()
            reports = parser.discover_reports(raw_root, "00123")
            result = parser.parse_reports(reports, "00123")

        dataset = result.dataset
        validate_dataset_for_tax_forms(dataset)
        self.assertEqual(dataset.warnings, [])
        self.assertEqual(dataset.tables["Trades"][0]["commission"], "2")
        self.assertEqual(len(dataset.tables["Transfers"]), 2)
        self.assertEqual(dataset.tables["Unprocessed"][0]["reason"], "unmatched_cash_redemption")
        self.assertEqual(dataset.tables["Unprocessed"][0]["isin"], "KZ2000000002")
        self.assertEqual(dataset.tables["Dividends"][0]["gross_amount"], "10.00")
        self.assertEqual(dataset.tables["Coupons"][0]["gross_amount"], "1.00")
        self.assertEqual(dataset.tables["CashBalances"][0]["ending_cash"], "141.00")
        self.assertEqual(dataset.tables["CorporateActions"][0]["proceeds"], "30.00")
        self.assertEqual(
            [row["year"] for row in dataset.tables["Years_Results"] if row["table"] == "Yearly Dividends"],
            [2024],
        )

        reconciliation = ReconciliationEngine().reconcile_dataset(dataset)
        self.assertEqual(
            [row for row in reconciliation if row.severity == ReconciliationSeverity.ERROR],
            [],
        )

    def test_xlsx_closing_positions_are_parsed(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "00123 2025 ОДЦБ.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            for row in (
                ("Отчет движения ценных бумаг",),
                (None, "№ лицевого счета:", "00123"),
                (None, "Отчет составлен на:", "01.01.2025-31.12.2025"),
                ("Ценные бумаги в портфеле клиента на конец отчетного периода:",),
                ("№", "Эмитент", "Вид ЦБ", "ISIN", "Номинал", "Количество ЦБ"),
                ("1", "ЭМИТЕНТ А", "АКЦИИ", "KZ0000000001", "1", 10),
            ):
                sheet.append(row)
            workbook.save(path)
            workbook.close()

            report = parse_alatay_xlsx(path)

        self.assertEqual(report.period_end.isoformat(), "2025-12-31")
        self.assertEqual(report.positions[0]["quantity"], "10")
        self.assertEqual(report.positions[0]["_snapshot_kind"], "closing")

    def test_xlsx_dash_nominal_is_treated_as_missing_value(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "00123 2025 ОДЦБ.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            for row in (
                ("Отчет движения ценных бумаг",),
                (None, "№ лицевого счета:", "00123"),
                (None, "Отчет составлен на:", "01.01.2025-31.12.2025"),
                ("Ценные бумаги в портфеле клиента на конец отчетного периода:",),
                ("№", "Эмитент", "Вид ЦБ", "ISIN", "Номинал", "Количество ЦБ"),
                ("1", "ЭМИТЕНТ А", "АКЦИИ", "KZ0000000001", "-", 0.0022),
            ):
                sheet.append(row)
            workbook.save(path)
            workbook.close()

            report = parse_alatay_xlsx(path)

        self.assertEqual(report.positions[0]["nominal"], "0")
        self.assertEqual(report.positions[0]["quantity"], "0.0022")

    def test_compact_xlsx_cash_sections_build_income_and_currency_balances(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "00123 2025 moneyMoves.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            for row in (
                ("ОТЧЕТ ДВИЖЕНИЯ ДЕНЕЖНЫХ СРЕДСТВ",),
                (None, "№ лицевого счета:", "00123"),
                (None, "Отчет составлен на:", "01.01.2025-31.12.2025"),
                ("Входящий остаток на начало периода:", "1.50 USD", "100 KZT"),
                ("Исходящий остаток на конец периода:", "2.50 USD", "110 KZT"),
                ("Движение денежных средств по счету в", "Долларах США"),
                (
                    "Дата проведения операции/сделки",
                    "Содержание операции/сделки",
                    "Тип операции/вид сделки",
                    "Эмитент/ISIN",
                    "Сумма Долларах США",
                ),
                (None, None, None, None, "Входящий остаток", "Приход", "Расход", "Исходящий остаток"),
                (
                    "24.06.2025",
                    "Зачисление денежных средств",
                    "Зачисление вознаграждения (дивиденды)",
                    "TEST INC/ US0000000001",
                    1.5,
                    1,
                    0,
                    2.5,
                ),
                ("Движение денежных средств по счету в", "Казахстанских тенге"),
                (
                    "Дата проведения операции/сделки",
                    "Содержание операции/сделки",
                    "Тип операции/вид сделки",
                    "Эмитент/ISIN",
                    "Сумма Казахстанских тенге",
                ),
                (None, None, None, None, "Входящий остаток", "Приход", "Расход", "Исходящий остаток"),
                (
                    "25.06.2025",
                    "Зачисление денежных средств",
                    "Зачисление вознаграждения (купон)",
                    "ТЕСТ/ KZ0000000001",
                    100,
                    10,
                    0,
                    110,
                ),
            ):
                sheet.append(row)
            workbook.save(path)
            workbook.close()

            report = parse_alatay_xlsx(path)

        self.assertEqual([row["currency"] for row in report.cash_movements], ["USD", "KZT"])
        self.assertEqual(report.ending_cash_by_currency, {"USD": Decimal("2.50"), "KZT": Decimal("110")})
        dataset = build_canonical_dataset(
            [report],
            "00123",
            AnnualFxRateProvider({(2025, "USD"): Decimal("500")}),
        )
        self.assertEqual(dataset.tables["Dividends"][0]["gross_amount"], "1.00")
        self.assertEqual(dataset.tables["Coupons"][0]["gross_amount"], "10.00")
        self.assertEqual(
            {(row["currency"], row["ending_cash"]) for row in dataset.tables["CashBalances"]},
            {("USD", "2.50"), ("KZT", "110.00")},
        )

    def test_redemption_is_a_closing_security_event(self) -> None:
        row = _parse_trade_row(
            [
                "20.11.2025",
                "ЭМИТЕНТ",
                "ОБЛИГАЦИИ",
                "KZ2C00010171",
                "1 000",
                "Погашение ЦБ",
                "57",
                "1 000",
                "KZT",
                "57 000",
                "KASE",
                "KZ",
                "0",
            ],
            source_report="report.xlsx",
            source_row=10,
        )

        self.assertIsNotNone(row)
        assert row is not None
        self.assertTrue(row["_recognized_operation"])
        self.assertEqual(row["quantity"], "-57")

    def test_bond_redemption_is_classified_in_trades_and_yearly_results(self) -> None:
        report = ParsedAlatayReport(
            path=Path("report.xlsx"),
            period_end=date(2025, 12, 31),
            trades=[
                {
                    "date_time": "2025-01-15 00:00:00",
                    "issuer": "Issuer",
                    "security_type": "Облигации",
                    "isin": "KZ2C00010171",
                    "operation": "Покупка",
                    "quantity": "57",
                    "price": "900",
                    "amount": "51300",
                    "currency": "KZT",
                    "exchange": "KASE",
                    "commission": "0",
                    "source_report": "report.xlsx",
                    "source_row": 10,
                },
                {
                    "date_time": "2025-11-20 00:00:00",
                    "issuer": "Issuer",
                    "security_type": "Облигации",
                    "isin": "KZ2C00010171",
                    "operation": "Погашение ЦБ",
                    "quantity": "-57",
                    "price": "1000",
                    "amount": "57000",
                    "currency": "KZT",
                    "exchange": "KASE",
                    "commission": "0",
                    "source_report": "report.xlsx",
                    "source_row": 11,
                },
            ],
        )

        dataset = build_canonical_dataset(
            [report],
            "ATEST",
            AnnualFxRateProvider({(2025, "KZT"): Decimal("1")}),
        )

        redemption = dataset.tables["Trades"][1]
        self.assertEqual(redemption["asset_type"], "Bonds")
        self.assertEqual(redemption["trade_type"], "corporate_action:redemption")
        self.assertEqual(dataset.tables["Fifo"][0]["operation_type"], "bond_redemption")
        self.assertEqual(dataset.tables["Fifo"][0]["years_result_table"], "Yearly Bonds Redemption")
        self.assertEqual(dataset.tables["Years_Results"][0]["table"], "Yearly Bonds Redemption")

    def test_cash_redemption_without_security_evidence_is_unprocessed(self) -> None:
        report = ParsedAlatayReport(
            path=Path("cash.csv"),
            period_end=date(2024, 12, 31),
            cash_movements=[
                {
                    "date": "2024-07-22",
                    "description": "Cash credited for a trade",
                    "operation": "Погашение ЦБ KZ2C00008654",
                    "issuer": "Issuer",
                    "isin": "KZ2C00008654",
                    "credit": "40000",
                    "debit": "0",
                    "currency": "KZT",
                    "security_type": "ОБЛИГАЦИИ",
                    "issuer_country": "KZ",
                    "source_report": "cash.csv",
                    "source_row": 17,
                }
            ],
        )

        dataset = build_canonical_dataset([report], "ATEST", AnnualFxRateProvider({}))

        self.assertEqual(dataset.tables["CorporateActions"][0]["proceeds"], "40000.00")
        self.assertEqual(len(dataset.tables["Unprocessed"]), 1)
        issue = dataset.tables["Unprocessed"][0]
        self.assertEqual(issue["reason"], "unmatched_cash_redemption")
        self.assertEqual(issue["severity"], "warning")
        self.assertEqual(issue["isin"], "KZ2C00008654")

        reconciliation = ReconciliationEngine().reconcile_dataset(dataset)
        unprocessed_rows = [row for row in reconciliation if row.metric == ReconciliationMetric.UNPROCESSED_ROWS]
        self.assertEqual(len(unprocessed_rows), 1)
        self.assertEqual(unprocessed_rows[0].severity, ReconciliationSeverity.WARNING)

    def test_cash_redemption_with_matching_security_redemption_is_processed(self) -> None:
        report = ParsedAlatayReport(
            path=Path("combined.xlsx"),
            period_end=date(2024, 12, 31),
            trades=[
                {
                    "date_time": "2024-01-15 00:00:00",
                    "issuer": "Issuer",
                    "security_type": "ОБЛИГАЦИИ",
                    "isin": "KZ2C00008654",
                    "operation": "Покупка",
                    "quantity": "40",
                    "price": "950",
                    "amount": "38000",
                    "currency": "KZT",
                    "exchange": "KASE",
                    "commission": "0",
                    "source_report": "combined.xlsx",
                    "source_row": 10,
                },
                {
                    "date_time": "2024-07-22 00:00:00",
                    "issuer": "Issuer",
                    "security_type": "ОБЛИГАЦИИ",
                    "isin": "KZ2C00008654",
                    "operation": "Погашение ЦБ",
                    "quantity": "-40",
                    "price": "1000",
                    "amount": "40000",
                    "currency": "KZT",
                    "exchange": "KASE",
                    "commission": "0",
                    "source_report": "combined.xlsx",
                    "source_row": 11,
                },
            ],
            cash_movements=[
                {
                    "date": "2024-07-22",
                    "description": "Cash credited for a trade",
                    "operation": "Погашение ЦБ KZ2C00008654",
                    "issuer": "Issuer",
                    "isin": "KZ2C00008654",
                    "credit": "40000",
                    "debit": "0",
                    "currency": "KZT",
                    "security_type": "ОБЛИГАЦИИ",
                    "issuer_country": "KZ",
                    "source_report": "cash.csv",
                    "source_row": 17,
                }
            ],
        )

        dataset = build_canonical_dataset([report], "ATEST", AnnualFxRateProvider({}))

        self.assertEqual(dataset.tables["Unprocessed"], [])

    def test_unexplained_earliest_position_becomes_missing_opening_lot(self) -> None:
        report = ParsedAlatayReport(
            path=Path("incomplete.xlsx"),
            account_id="00123",
            period_start=date(2025, 1, 1),
            period_end=date(2025, 12, 31),
            positions=[
                {
                    "issuer": "ЭМИТЕНТ",
                    "security_type": "АКЦИИ",
                    "isin": "KZ0000000001",
                    "nominal": "1",
                    "quantity": "10",
                    "issuer_country": "KZ",
                    "source_report": "incomplete.xlsx",
                    "source_row": 10,
                    "_snapshot_kind": "closing",
                    "_snapshot_date": "2025-12-31",
                }
            ],
        )

        dataset = build_canonical_dataset([report], "00123", AnnualFxRateProvider({}))
        self.assertEqual(dataset.tables["Positions"][0]["quantity"], "10")
        self.assertEqual(dataset.tables["Unprocessed"][0]["reason"], "missing_opening_lot")
        reconciliation = ReconciliationEngine().reconcile_dataset(dataset)
        ending_rows = [row for row in reconciliation if row.metric == ReconciliationMetric.ENDING_POSITION_QUANTITY]
        self.assertTrue(all(row.severity == ReconciliationSeverity.INFO for row in ending_rows))

    def test_ownership_transfer_explains_earliest_position(self) -> None:
        report = ParsedAlatayReport(
            path=Path("transfer.xlsx"),
            account_id="00123",
            period_start=date(2023, 1, 1),
            period_end=date(2023, 12, 31),
            positions=[
                {
                    "issuer": "ЭМИТЕНТ",
                    "security_type": "АКЦИИ",
                    "isin": "KZ0000000001",
                    "nominal": "1",
                    "quantity": "92",
                    "issuer_country": "KZ",
                    "source_report": "transfer.xlsx",
                    "source_row": 10,
                    "_snapshot_kind": "closing",
                    "_snapshot_date": "2023-12-31",
                }
            ],
            security_transfers=[
                {
                    "date_time": "2023-05-11 00:00:00",
                    "operation": "Перевод основной (получатель смена прав собственности)",
                    "security_type": "АКЦИИ",
                    "isin": "KZ0000000001",
                    "quantity": "92",
                    "price": "158.9",
                    "currency": "KZT",
                    "issuer_country": "KZ",
                    "exchange": "KASE",
                    "source_report": "transfer.xlsx",
                }
            ],
        )

        dataset = build_canonical_dataset([report], "00123", AnnualFxRateProvider({}))
        self.assertEqual(dataset.tables["Positions"][0]["quantity"], "92")
        self.assertEqual(dataset.tables["Unprocessed"], [])
        reconciliation = ReconciliationEngine().reconcile_dataset(dataset)
        ending_rows = [row for row in reconciliation if row.metric == ReconciliationMetric.ENDING_POSITION_QUANTITY]
        self.assertTrue(all(row.severity == ReconciliationSeverity.INFO for row in ending_rows))

    def test_registry_exposes_alatay(self) -> None:
        self.assertIn("alatay", default_registry().broker_codes())

    def test_account_detection_restores_leading_zero_from_alatay_filename(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "01010105826 report.csv"
            path.write_text(REPORT_2023.replace("ATEST", "1010105826"), encoding="utf-8")

            self.assertEqual(detect_report_account_id("alatay", path), "01010105826")


if __name__ == "__main__":
    unittest.main()
