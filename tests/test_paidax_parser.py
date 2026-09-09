from __future__ import annotations

import tempfile
import unittest
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook

from conftest_imports import SRC  # noqa: F401
from kztax270.brokers.account_detection import detect_report_account_id, detect_report_period_end
from kztax270.brokers.paidax import PaidaxParser, parse_paidax_xlsx
from kztax270.brokers.registry import default_registry
from kztax270.canonical.validation import validate_dataset_for_tax_forms
from kztax270.reconciliation.engine import ReconciliationEngine
from kztax270.reconciliation.models import ReconciliationSeverity
from kztax270.reference.fx import AnnualFxRateProvider

TRADE_HEADERS = (
    "Дата",
    "Тикер",
    "Наименование",
    "ISIN",
    "Код страны регистрации эмитента",
    "Наименование рынка",
    "Тип актива",
    "Метод торгов",
    "Тип операции",
    "Кол-во",
    "Цена",
    "Сумма сделки",
    "Валюта",
    "Комиссия брокера",
    "Биржевой сбор",
    "Итого комиссии",
    "Прибыль / убыток (валюта сделки)",
    "Курс НБ РК",
    "Прибыль / убыток (KZT)",
)


def _report(path: Path, account_id: str = "PDK001") -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Сводка"
    summary.cell(
        3,
        2,
        f"Период: 01.01.2025 — 31.12.2025 (Asia/Qyzylorda) Клиент: Test Договор № {account_id} от 01.01.2025",
    )

    trades = workbook.create_sheet("Сделки")
    for column, header in enumerate(TRADE_HEADERS, start=2):
        trades.cell(4, column, header)
    trades.append(
        (
            None,
            "01.02.2025",
            "AAA",
            "AAA Inc.",
            "US0000000001",
            "USA",
            "ITS",
            "Акция",
            "Биржевой",
            "Покупка",
            2,
            10,
            20,
            "USD",
            0.8,
            0.2,
            1,
            "—",
            520,
            "—",
        )
    )
    trades.append(
        (
            None,
            "01.03.2025",
            "AAA",
            "AAA Inc.",
            "US0000000001",
            "USA",
            "ITS",
            "Акция",
            "Биржевой",
            "Продажа",
            1,
            15,
            15,
            "USD",
            0.8,
            0.2,
            1,
            4,
            520,
            2080,
        )
    )
    trades.append(
        (
            None,
            "01.04.2025",
            "BTC/USD",
            "Bitcoin",
            None,
            None,
            "Crypto",
            "Криптовалюта",
            "Биржевой",
            "Покупка",
            0.1,
            60000,
            6000,
            "USD",
            1,
            0,
            1,
            "—",
            520,
            "—",
        )
    )

    positions = workbook.create_sheet("Позиции")
    position_headers = (
        "Тикер",
        "Наименование",
        "ISIN",
        "Код страны эмитента",
        "Тип актива",
        "Адрес регистрации",
        "Остаток нач.",
        "Остаток кон.",
        "Цена на конец",
        "Оценка на конец",
        "Валюта",
    )
    for column, header in enumerate(position_headers, start=2):
        positions.cell(3, column, header)
    positions.append((None, "AAA", "AAA Inc.", "US0000000001", "USA", "Акция", None, 0, 1, 15, 15, "USD"))
    positions.append((None, "BTC/USD", "Bitcoin", None, None, "Криптовалюта", None, 0, 0.1, 60000, 6000, "USD"))

    payments = workbook.create_sheet("Выплаты")
    payment_headers = (
        "Дата выплаты",
        "Тип выплаты",
        "Тикер",
        "ISIN",
        "Наименование",
        "Наименование рынка",
        "Код страны регистрации эмитента",
        "Тип ценной бумаги",
        "Дата фиксации",
        "Кол-во на фикс.",
        "На 1 бумагу",
        "Валюта",
        "Начислено (gross)",
        "Налог у источника",
        "Получено (net)",
    )
    for column, header in enumerate(payment_headers, start=2):
        payments.cell(4, column, header)
    payments.append((None, "15.06.2025", "Дивиденд", "AAA", "US0000000001", "AAA Inc.", "ITS", "USA", "Акция", "10.06.2025", 1, 2, "USD", 2, 0.3, 1.7))

    cash = workbook.create_sheet("Деньги")
    for column, header in enumerate(
        ("Валюта", "Остаток на начало периода", "Остаток на конец периода"), start=2
    ):
        cash.cell(5, column, header)
    cash.append((None, "USD", 0, 100))
    cash.append((None, "ИТОГО", 0, 100))

    workbook.save(path)
    workbook.close()


class PaidaxParserTests(unittest.TestCase):
    def test_paidax_builds_canonical_audit_and_excludes_direct_crypto(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            broker_root = root / "paidax"
            broker_root.mkdir()
            path = broker_root / "PDK001 report.xlsx"
            _report(path)
            parser = PaidaxParser(
                AnnualFxRateProvider({(2025, "USD"): Decimal("520")})
            )
            result = parser.parse_reports(parser.discover_reports(root, "PDK001"), "PDK001")

        dataset = result.dataset
        validate_dataset_for_tax_forms(dataset)
        self.assertEqual(dataset.metadata.broker, "paidax")
        self.assertEqual([row["symbol"] for row in dataset.tables["Trades"]], ["AAA", "AAA"])
        self.assertEqual(len(dataset.tables["Fifo"]), 1)
        self.assertEqual(dataset.tables["Dividends"][0]["withholding_tax"], "-0.30")
        self.assertEqual(dataset.tables["CashBalances"][0]["ending_cash"], "100.00")
        crypto = [
            row
            for row in dataset.tables["Unprocessed"]
            if row["reason"] == "unsupported_crypto_asset"
        ]
        self.assertEqual({row["source_sheet"] for row in crypto}, {"Сделки", "Позиции"})
        self.assertTrue(all(row.get("symbol") == "BTC/USD" for row in crypto))
        self.assertNotIn("BTC/USD", {row["symbol"] for row in dataset.tables["Instruments"]})
        reconciliation = ReconciliationEngine().reconcile_dataset(dataset)
        self.assertEqual(
            [row for row in reconciliation if row.severity == ReconciliationSeverity.ERROR],
            [],
        )

    def test_account_detection_and_registry_expose_paidax(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "report.xlsx"
            _report(path, "PDK008156")
            parsed = parse_paidax_xlsx(path)
            self.assertEqual(parsed.account_id, "PDK008156")
            self.assertEqual(detect_report_account_id("paidax", path), "PDK008156")
            self.assertEqual(detect_report_period_end("paidax", path).isoformat(), "2025-12-31")
        self.assertIn("paidax", default_registry().broker_codes())


if __name__ == "__main__":
    unittest.main()
