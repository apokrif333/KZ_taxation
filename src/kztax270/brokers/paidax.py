"""Parser for Paidax brokerage report workbooks."""

from __future__ import annotations

import re
from collections.abc import Mapping, Sequence
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from kztax270.canonical.schema import AccountMetadata, CanonicalDataset
from kztax270.canonical.trade_enrichment import enrich_trades_before_calculations
from kztax270.diagnostics import instrument_parsed_reports
from kztax270.reconciliation.models import ReconciliationMetric
from kztax270.reference.fx import AnnualFxRateProvider

from .base import BrokerReport, ParseResult
from .discovery import DiscoveryRule, discover_raw_reports
from .ib import (
    FifoOpenLot,
    _build_broker_trade_realized_pl,
    _build_fifo_and_positions,
    _build_unprocessed_rows,
    _build_years_results,
    _canonical_trade_rows,
    _canonical_transfer_rows,
    _instrument_identity_key_from_values,
    _instrument_symbol_history,
    _sort_trades_by_datetime,
)

BROKER_CODE = "paidax"
RAW_FOLDER = "paidax"
BASE_CURRENCY = "KZT"

SUMMARY_SHEET = "Сводка"
POSITIONS_SHEET = "Позиции"
TRADES_SHEET = "Сделки"
SECURITY_MOVEMENTS_SHEET = "Движение бумаг"
CORPORATE_ACTIONS_SHEET = "Корп. события"
PAYMENTS_SHEET = "Выплаты"
SAVINGS_SHEET = "Сбережения"
CASH_SHEET = "Деньги"

_DATE_PATTERN = re.compile(r"\d{2}\.\d{2}\.\d{4}")
_ACCOUNT_PATTERN = re.compile(r"Договор\s*№\s*([^\s]+)", re.IGNORECASE)
_DIRECT_CRYPTO_SYMBOLS = {
    "ADA",
    "BNB",
    "BTC",
    "DOGE",
    "DOT",
    "ETH",
    "LTC",
    "SOL",
    "TRX",
    "USDC",
    "USDT",
    "XRP",
}


@dataclass(slots=True)
class ParsedPaidaxReport:
    path: Path
    account_id: str | None = None
    period_start: date | None = None
    period_end: date | None = None
    trades: list[dict[str, Any]] = field(default_factory=list)
    security_movements: list[dict[str, Any]] = field(default_factory=list)
    corporate_actions: list[dict[str, Any]] = field(default_factory=list)
    payments: list[dict[str, Any]] = field(default_factory=list)
    positions: list[dict[str, Any]] = field(default_factory=list)
    savings: list[dict[str, Any]] = field(default_factory=list)
    cash_balances: list[dict[str, Any]] = field(default_factory=list)
    cash_transfers: list[dict[str, Any]] = field(default_factory=list)
    unprocessed: list[dict[str, Any]] = field(default_factory=list)


class PaidaxParser:
    broker_code = BROKER_CODE

    def __init__(self, fx_provider: AnnualFxRateProvider | None = None) -> None:
        self.fx_provider = fx_provider or AnnualFxRateProvider({})

    def discover_reports(self, raw_root: Path, account_id: str) -> list[BrokerReport]:
        return discover_raw_reports(
            raw_root,
            DiscoveryRule(
                broker=RAW_FOLDER,
                account_id=account_id,
                extensions=frozenset({".xlsx"}),
            ),
        )

    def parse_reports(self, reports: Sequence[BrokerReport], account_id: str) -> ParseResult:
        parsed_reports = [parse_paidax_xlsx(report.path) for report in reports]
        instrument_parsed_reports(self.broker_code, parsed_reports)
        dataset = build_canonical_dataset(parsed_reports, account_id, self.fx_provider)
        dataset.raw_totals.source_reports = [str(report.path) for report in reports]
        return ParseResult(
            broker=self.broker_code,
            account_id=account_id,
            reports=reports,
            dataset=dataset,
            raw_totals=dataset.raw_totals,
        )


def parse_paidax_xlsx(path: Path) -> ParsedPaidaxReport:
    workbook = load_workbook(path, read_only=True, data_only=True)
    parsed = ParsedPaidaxReport(path=path)
    try:
        # Paidax workbooks can contain stale worksheet ``<dimension>`` values.
        # In read-only mode openpyxl treats that value as an iteration boundary,
        # which otherwise silently drops valid rows below the declared range.
        for worksheet in workbook.worksheets:
            worksheet.reset_dimensions()
            worksheet.calculate_dimension(force=True)
        if SUMMARY_SHEET in workbook.sheetnames:
            _parse_summary(workbook[SUMMARY_SHEET], parsed)
        if TRADES_SHEET in workbook.sheetnames:
            _parse_trades(workbook[TRADES_SHEET], parsed)
        if SECURITY_MOVEMENTS_SHEET in workbook.sheetnames:
            _parse_security_movements(workbook[SECURITY_MOVEMENTS_SHEET], parsed)
        if CORPORATE_ACTIONS_SHEET in workbook.sheetnames:
            _parse_corporate_actions(workbook[CORPORATE_ACTIONS_SHEET], parsed)
        if PAYMENTS_SHEET in workbook.sheetnames:
            _parse_payments(workbook[PAYMENTS_SHEET], parsed)
        if POSITIONS_SHEET in workbook.sheetnames:
            _parse_positions(workbook[POSITIONS_SHEET], parsed)
        if SAVINGS_SHEET in workbook.sheetnames:
            _parse_savings(workbook[SAVINGS_SHEET], parsed)
        if CASH_SHEET in workbook.sheetnames:
            _parse_cash(workbook[CASH_SHEET], parsed)
    finally:
        workbook.close()
    return parsed


def build_canonical_dataset(
    reports: Sequence[ParsedPaidaxReport],
    account_id: str,
    fx_provider: AnnualFxRateProvider,
) -> CanonicalDataset:
    dataset = CanonicalDataset(
        metadata=AccountMetadata(broker=BROKER_CODE, account_id=account_id, base_currency=BASE_CURRENCY)
    )
    for report in reports:
        if report.account_id and report.account_id.casefold() != account_id.casefold():
            dataset.warnings.append(
                f"Paidax report {report.path} belongs to account {report.account_id}, expected {account_id}."
            )

    instruments = _build_instruments(reports, account_id)
    dataset.tables["Instruments"] = instruments
    internal_trades = _sort_trades_by_datetime(
        [
            *_build_trades(reports, instruments),
            *_build_gift_trades(reports, instruments),
            *_build_corporate_action_trades(reports, instruments),
        ]
    )
    enrich_trades_before_calculations(dataset, internal_trades, fx_provider)
    transfers = _build_transfers(reports, instruments)
    dataset.tables["Trades"] = _canonical_trade_rows(internal_trades)
    dataset.tables["_BrokerTradeRealizedPL"] = _build_broker_trade_realized_pl(internal_trades)

    fifo_rows, positions, transfer_rows = _build_fifo_and_positions(
        internal_trades,
        transfers=transfers,
        initial_lots=_build_initial_lots(reports, instruments),
        max_year=_max_report_year(reports),
        fx_provider=fx_provider,
        warnings=dataset.warnings,
        symbol_history=_instrument_symbol_history(instruments),
        join_security_currencies=True,
    )
    dataset.tables["Fifo"] = fifo_rows
    dataset.tables["Positions"] = positions
    dataset.tables["Transfers"] = _canonical_transfer_rows(transfer_rows)
    dataset.tables["CorporateActions"] = _build_corporate_actions(reports, instruments)
    dataset.tables["Dividends"] = _build_dividends(reports, instruments, fx_provider, dataset.warnings)
    dataset.tables["Interest"] = _build_interest(reports, fx_provider, dataset.warnings)
    dataset.tables["Coupons"] = []
    dataset.tables["CashBalances"] = _build_cash_balances(
        reports, account_id, fx_provider, dataset.warnings
    )
    dataset.tables["Unprocessed"] = [
        *_build_unprocessed_rows(dataset.tables["Trades"], fifo_rows),
        *(row for report in reports for row in report.unprocessed),
    ]
    dataset.tables["Years_Results"] = _build_years_results(dataset)
    _populate_raw_totals(dataset, reports, internal_trades)
    return dataset


def _parse_summary(worksheet: Any, parsed: ParsedPaidaxReport) -> None:
    text = " ".join(
        _clean_text(cell.value)
        for row in worksheet.iter_rows(min_row=1, max_row=min(8, worksheet.max_row))
        for cell in row
        if _clean_text(cell.value)
    )
    account_match = _ACCOUNT_PATTERN.search(text)
    if account_match:
        parsed.account_id = account_match.group(1).strip()
    dates = [_parse_date(value) for value in _DATE_PATTERN.findall(text)]
    valid_dates = [value for value in dates if value is not None]
    if len(valid_dates) >= 2:
        parsed.period_start = valid_dates[0]
        parsed.period_end = valid_dates[1]


def _parse_trades(worksheet: Any, parsed: ParsedPaidaxReport) -> None:
    for source_row, row in _records(worksheet, ("Дата", "Тикер", "Тип операции")):
        if _parse_date(row.get("Дата")) is None:
            continue
        normalized = _security_row(row, parsed.path, source_row, TRADES_SHEET)
        if _is_crypto(normalized):
            parsed.unprocessed.append(_crypto_unprocessed(normalized))
            continue
        operation = _text(row.get("Тип операции"))
        if not operation or operation.casefold() not in {"покупка", "продажа"}:
            parsed.unprocessed.append(
                _unprocessed(normalized, "unsupported_operation", f"Unsupported Paidax trade operation: {operation}")
            )
            continue
        normalized.update(
            {
                "date_time": _datetime_text(row.get("Дата")),
                "operation": operation,
                "quantity": str(_decimal(row.get("Кол-во"))),
                "price": str(_decimal(row.get("Цена"))),
                "amount": str(abs(_decimal(row.get("Сумма сделки")))),
                "commission": str(abs(_decimal(row.get("Итого комиссии")))),
                "broker_realized_pl": _optional_decimal_text(
                    row.get("Прибыль / убыток (валюта сделки)")
                ),
            }
        )
        parsed.trades.append(normalized)


def _parse_security_movements(worksheet: Any, parsed: ParsedPaidaxReport) -> None:
    for source_row, row in _records(worksheet, ("Дата", "Тикер", "Тип операции")):
        if _parse_date(row.get("Дата")) is None:
            continue
        normalized = _security_row(row, parsed.path, source_row, SECURITY_MOVEMENTS_SHEET)
        normalized.update(
            {
                "date_time": _datetime_text(row.get("Дата")),
                "operation": _text(row.get("Тип операции")),
                "method": _text(row.get("Метод торгов")),
                "direction": _text(row.get("Направл. (+/−)")),
                "quantity": str(abs(_decimal(row.get("Кол-во")))),
                "price": str(_decimal(row.get("Цена на дату (рын.)"))),
                "amount": str(abs(_decimal(row.get("Оценочная стоимость")))),
                "acquisition_price": str(_decimal(row.get("Цена приобретения"))),
                "broker_realized_pl": _optional_decimal_text(row.get("Прибыль / убыток (валюта)")),
            }
        )
        if _is_crypto(normalized):
            parsed.unprocessed.append(_crypto_unprocessed(normalized))
            continue
        parsed.security_movements.append(normalized)


def _parse_corporate_actions(worksheet: Any, parsed: ParsedPaidaxReport) -> None:
    for source_row, row in _records(worksheet, ("Дата события", "Тип события", "Тикер")):
        event_date = _parse_date(row.get("Дата события"))
        if event_date is None:
            continue
        normalized = _security_row(row, parsed.path, source_row, CORPORATE_ACTIONS_SHEET)
        normalized.update(
            {
                "date": event_date.isoformat(),
                "action_type": _text(row.get("Тип события")),
                "record_date": _date_text(row.get("Дата фиксации")),
                "record_quantity": str(_decimal(row.get("Бумаг на фикс."))),
                "before_quantity": str(_decimal(row.get("До события"))),
                "after_quantity": str(_decimal(row.get("После события"))),
                "parameter": _text(row.get("Коэф./ Параметр")),
                "received": str(_decimal(row.get("Получено"))),
                "comment": _text(row.get("Комментарий")),
            }
        )
        if _is_crypto(normalized):
            parsed.unprocessed.append(_crypto_unprocessed(normalized))
            continue
        parsed.corporate_actions.append(normalized)


def _parse_payments(worksheet: Any, parsed: ParsedPaidaxReport) -> None:
    for source_row, row in _records(worksheet, ("Дата выплаты", "Тип выплаты", "Тикер")):
        paid_at = _parse_date(row.get("Дата выплаты"))
        if paid_at is None:
            continue
        normalized = _security_row(row, parsed.path, source_row, PAYMENTS_SHEET)
        normalized.update(
            {
                "date": paid_at.isoformat(),
                "payment_type": _text(row.get("Тип выплаты")),
                "record_date": _date_text(row.get("Дата фиксации")),
                "quantity": str(_decimal(row.get("Кол-во на фикс."))),
                "per_security": str(_decimal(row.get("На 1 бумагу"))),
                "gross": str(_decimal(row.get("Начислено (gross)"))),
                "withholding": str(abs(_decimal(row.get("Налог у источника")))),
                "net": str(_decimal(row.get("Получено (net)"))),
            }
        )
        if _is_crypto(normalized):
            parsed.unprocessed.append(_crypto_unprocessed(normalized))
            continue
        if "дивиденд" not in str(normalized.get("payment_type") or "").casefold():
            parsed.unprocessed.append(
                _unprocessed(
                    normalized,
                    "unsupported_income",
                    f"Unsupported Paidax payment type: {normalized.get('payment_type')}",
                )
            )
            continue
        parsed.payments.append(normalized)


def _parse_positions(worksheet: Any, parsed: ParsedPaidaxReport) -> None:
    for source_row, row in _records(worksheet, ("Тикер", "ISIN", "Остаток кон.")):
        symbol = _text(row.get("Тикер"))
        if not symbol or symbol.casefold().startswith("итого"):
            continue
        normalized = _security_row(row, parsed.path, source_row, POSITIONS_SHEET)
        normalized.update(
            {
                "quantity": str(_decimal(row.get("Остаток кон."))),
                "starting_quantity": str(_decimal(row.get("Остаток нач."))),
                "price": str(_decimal(row.get("Цена на конец"))),
                "amount": str(_decimal(row.get("Оценка на конец"))),
            }
        )
        if _is_crypto(normalized):
            parsed.unprocessed.append(_crypto_unprocessed(normalized))
            continue
        parsed.positions.append(normalized)


def _parse_savings(worksheet: Any, parsed: ParsedPaidaxReport) -> None:
    for source_row, row in _records(worksheet, ("Инструмент / Продукт", "Остаток (кон.)", "Валюта")):
        product = _text(row.get("Инструмент / Продукт"))
        if not product or product.casefold().startswith("итого"):
            continue
        parsed.savings.append(
            {
                "date": parsed.period_end.isoformat() if parsed.period_end else None,
                "product": product,
                "savings_type": _text(row.get("Тип")),
                "starting_balance": str(_decimal(row.get("Остаток (нач.)"))),
                "ending_balance": str(_decimal(row.get("Остаток (кон.)"))),
                "income": str(_decimal(row.get("Начислено дохода"))),
                "currency": _currency(row.get("Валюта")) or BASE_CURRENCY,
                "source_report": str(parsed.path),
                "source_row": source_row,
                "source_sheet": SAVINGS_SHEET,
            }
        )


def _parse_cash(worksheet: Any, parsed: ParsedPaidaxReport) -> None:
    summary_header = _find_header_row(
        worksheet, ("Валюта", "Остаток на начало периода", "Остаток на конец периода")
    )
    if summary_header:
        headers = _headers(worksheet, summary_header)
        for row_index, values in enumerate(
            worksheet.iter_rows(min_row=summary_header + 1, values_only=True),
            start=summary_header + 1,
        ):
            row = _row_mapping(headers, values)
            currency = _currency(row.get("Валюта"))
            if currency in {"ИТОГО", "TOTAL"}:
                break
            if not currency or not re.fullmatch(r"[A-Z]{3}", currency):
                continue
            if _parse_date(row.get("Дата")) is not None:
                break
            parsed.cash_balances.append(
                {
                    "currency": currency,
                    "starting_cash": str(_decimal(row.get("Остаток на начало периода"))),
                    "ending_cash": str(_decimal(row.get("Остаток на конец периода"))),
                    "source_report": str(parsed.path),
                    "source_row": row_index,
                }
            )

    transfer_header = _find_header_row(worksheet, ("Дата", "Тип операции", "Валюта"))
    if transfer_header:
        headers = _headers(worksheet, transfer_header)
        for row_index, values in enumerate(
            worksheet.iter_rows(min_row=transfer_header + 1, values_only=True),
            start=transfer_header + 1,
        ):
            row = _row_mapping(headers, values)
            transfer_date = _parse_date(row.get("Дата"))
            if transfer_date is None:
                continue
            deposit = abs(_decimal(row.get("Сумма пополнения")))
            withdrawal = abs(_decimal(row.get("Сумма вывода")))
            amount = deposit - withdrawal
            if amount == 0:
                continue
            parsed.cash_transfers.append(
                {
                    "date": transfer_date.isoformat(),
                    "operation": _text(row.get("Тип операции")),
                    "currency": _currency(row.get("Валюта")) or BASE_CURRENCY,
                    "amount": str(amount),
                    "source_report": str(parsed.path),
                    "source_row": row_index,
                }
            )


def _build_instruments(
    reports: Sequence[ParsedPaidaxReport], account_id: str
) -> list[dict[str, Any]]:
    sources: dict[str, Mapping[str, Any]] = {}
    report_dates: dict[str, date | None] = {}
    for report in reports:
        rows = [
            *report.trades,
            *report.security_movements,
            *report.corporate_actions,
            *report.payments,
            *report.positions,
        ]
        for row in rows:
            isin = _text(row.get("isin"))
            if not isin:
                continue
            existing = sources.get(isin)
            if existing is None or (not _text(existing.get("exchange")) and _text(row.get("exchange"))):
                sources[isin] = row
            report_dates[isin] = _later_date(report_dates.get(isin), report.period_end)

    result: list[dict[str, Any]] = []
    for isin in sorted(sources):
        row = sources[isin]
        country = _country(row.get("issuer_country"), isin)
        result.append(
            {
                "symbol": _text(row.get("symbol")) or isin,
                "description": _text(row.get("name")) or isin,
                "conid": None,
                "security_id": isin,
                "underlying": None,
                "listing_exchange": _exchange(row.get("exchange")),
                "multiplier": "1",
                "type": _asset_type(row.get("security_type")),
                "code": None,
                "year": None,
                "expiry": None,
                "delivery_month": None,
                "strike": None,
                "issuer": _text(row.get("name")),
                "maturity": None,
                "cusip": None,
                "country": country,
                "isin": isin,
                "figi": None,
                "issuer_country": country,
                "offshore_flag": False if country == "KZ" else None,
                "issuer_outside_kz_flag": False if country == "KZ" else (True if country else None),
                "preferential_tax_flag": None,
                "source_broker": BROKER_CODE,
                "source_account": account_id,
                "source_report": row.get("source_report"),
                "as_of_date": report_dates[isin].isoformat() if report_dates.get(isin) else None,
            }
        )
    return result


def _build_trades(
    reports: Sequence[ParsedPaidaxReport], instruments: Sequence[Mapping[str, Any]]
) -> list[dict[str, Any]]:
    lookup = {str(row.get("isin")): row for row in instruments}
    result: list[dict[str, Any]] = []
    for report in reports:
        for row in report.trades:
            operation = str(row.get("operation") or "").casefold()
            raw_quantity = abs(_decimal(row.get("quantity")))
            quantity = -raw_quantity if operation == "продажа" else raw_quantity
            result.append(_canonical_internal_trade(report, row, lookup, quantity, "trade"))
    return result


def _build_initial_lots(
    reports: Sequence[ParsedPaidaxReport], instruments: Sequence[Mapping[str, Any]]
) -> list[tuple[tuple[str, str | None, str], str, FifoOpenLot]]:
    lookup = {str(row.get("isin")): row for row in instruments}
    dated_reports = [report for report in reports if report.period_start is not None]
    earliest_start = min((report.period_start for report in dated_reports), default=None)
    earliest_reports = [
        report for report in reports if earliest_start is None or report.period_start == earliest_start
    ]
    lots: list[tuple[tuple[str, str | None, str], str, FifoOpenLot]] = []
    seeded: dict[str, Decimal] = {}

    for report in earliest_reports:
        lot_date = datetime.combine(report.period_start, datetime.min.time()) if report.period_start else None
        for row in report.positions:
            signed_quantity = _decimal(row.get("starting_quantity"))
            if signed_quantity == 0:
                continue
            isin = _text(row.get("isin"))
            instrument = lookup.get(isin or "", {})
            symbol = _text(instrument.get("symbol")) or _text(row.get("symbol")) or isin or "UNKNOWN"
            currency = _currency(row.get("currency")) or BASE_CURRENCY
            quantity = abs(signed_quantity)
            key = _instrument_identity_key_from_values(isin=isin, symbol=symbol)
            if not key:
                continue
            lots.append(
                (
                    (key, isin, currency),
                    "long" if signed_quantity > 0 else "short",
                    FifoOpenLot(
                        asset_type=_text(instrument.get("type")) or _asset_type(row.get("security_type")),
                        symbol=symbol,
                        isin=isin,
                        currency=currency,
                        country=_text(instrument.get("country")) or _country(row.get("issuer_country"), isin),
                        exchange=_exchange(instrument.get("listing_exchange") or row.get("exchange")),
                        date_time=lot_date,
                        raw_quantity=signed_quantity,
                        raw_amount=Decimal("0"),
                        raw_commission=Decimal("0"),
                        price=Decimal("0"),
                        calculation_price=Decimal("0"),
                        multiplier=Decimal("1"),
                        quantity=quantity,
                        broker_quantity=quantity,
                        commission_per_unit=Decimal("0"),
                        trade_id=None,
                        opening_lot_status="missing_opening_lot",
                    ),
                )
            )
            if isin:
                seeded[isin] = seeded.get(isin, Decimal("0")) + signed_quantity

    for report in reports:
        ordinary_by_isin: dict[str, Decimal] = dict(seeded)
        timeline = sorted(
            [*report.trades, *report.security_movements],
            key=lambda row: str(row.get("date_time") or ""),
        )
        for row in timeline:
            isin = _text(row.get("isin"))
            if not isin:
                continue
            operation = str(row.get("operation") or "").casefold()
            quantity = abs(_decimal(row.get("quantity")))
            if row.get("source_sheet") == TRADES_SHEET:
                ordinary_by_isin[isin] = ordinary_by_isin.get(isin, Decimal("0")) + (
                    -quantity if operation == "продажа" else quantity
                )
                continue
            if "безвозмезд" not in operation and "подар" not in operation:
                continue
            if not _is_minus(row.get("direction")):
                ordinary_by_isin[isin] = ordinary_by_isin.get(isin, Decimal("0")) + quantity
                continue
            available = max(ordinary_by_isin.get(isin, Decimal("0")), Decimal("0"))
            shortage = max(quantity - available, Decimal("0"))
            ordinary_by_isin[isin] = available - quantity
            if shortage == 0:
                continue
            instrument = lookup.get(isin, {})
            symbol = _text(instrument.get("symbol")) or _text(row.get("symbol")) or isin
            currency = _currency(row.get("currency")) or BASE_CURRENCY
            price = abs(_decimal(row.get("acquisition_price")))
            event_dt = _parse_datetime(row.get("date_time"))
            lot_dt = event_dt - timedelta(seconds=1) if event_dt else None
            key = _instrument_identity_key_from_values(isin=isin, symbol=symbol)
            if not key:
                continue
            lots.append(
                (
                    (key, isin, currency),
                    "long",
                    FifoOpenLot(
                        asset_type=_text(instrument.get("type")) or _asset_type(row.get("security_type")),
                        symbol=symbol,
                        isin=isin,
                        currency=currency,
                        country=_text(instrument.get("country")) or _country(row.get("issuer_country"), isin),
                        exchange=_exchange(instrument.get("listing_exchange") or row.get("exchange")),
                        date_time=lot_dt,
                        raw_quantity=shortage,
                        raw_amount=shortage * price,
                        raw_commission=Decimal("0"),
                        price=price,
                        calculation_price=price,
                        multiplier=Decimal("1"),
                        quantity=shortage,
                        broker_quantity=shortage,
                        commission_per_unit=Decimal("0"),
                        trade_id=None,
                        opening_lot_status="broker_reported_cost_basis",
                    ),
                )
            )
    return lots


def _build_gift_trades(
    reports: Sequence[ParsedPaidaxReport], instruments: Sequence[Mapping[str, Any]]
) -> list[dict[str, Any]]:
    lookup = {str(row.get("isin")): row for row in instruments}
    result: list[dict[str, Any]] = []
    for report in reports:
        for row in report.security_movements:
            operation = str(row.get("operation") or "").casefold()
            if "безвозмезд" not in operation and "подар" not in operation:
                continue
            direction = str(row.get("direction") or "")
            raw_quantity = abs(_decimal(row.get("quantity")))
            quantity = -raw_quantity if _is_minus(direction) else raw_quantity
            result.append(
                _canonical_internal_trade(
                    report,
                    row,
                    lookup,
                    quantity,
                    "stock_award_withholding" if quantity < 0 else "stock_award_grant",
                )
            )
    return result


def _build_corporate_action_trades(
    reports: Sequence[ParsedPaidaxReport], instruments: Sequence[Mapping[str, Any]]
) -> list[dict[str, Any]]:
    lookup = {str(row.get("isin")): row for row in instruments}
    result: list[dict[str, Any]] = []
    for report in reports:
        prior_by_isin: dict[str, Mapping[str, Any]] = {}
        for row in sorted(report.corporate_actions, key=lambda item: str(item.get("date") or "")):
            isin = _text(row.get("isin"))
            action = str(row.get("action_type") or "").casefold()
            if isin and ("сплит" in action or "консолидац" in action):
                prior_by_isin[isin] = row
                continue
            if not isin or "денежн" not in action or "компенсац" not in action:
                continue
            received = abs(_decimal(row.get("received")))
            prior = prior_by_isin.get(isin)
            before = abs(_decimal(prior.get("before_quantity"))) if prior else Decimal("0")
            after = abs(_decimal(prior.get("after_quantity"))) if prior else Decimal("0")
            quantity = before - after
            if quantity <= 0 or received <= 0:
                report.unprocessed.append(
                    _unprocessed(
                        row,
                        "unsupported_corporate_action",
                        "Paidax cash compensation could not be matched to a disposed security quantity.",
                    )
                )
                continue
            synthetic = dict(row)
            synthetic.update(
                {
                    "quantity": str(quantity),
                    "price": str(received / quantity),
                    "amount": str(received),
                    "commission": "0",
                    "broker_realized_pl": None,
                    "date_time": f"{row.get('date')} 00:00:00",
                }
            )
            trade = _canonical_internal_trade(
                report, synthetic, lookup, -quantity, "corporate_action:split_compensation"
            )
            trade["_corporate_action_type"] = "split_compensation"
            trade["corporate_action_type"] = "split_compensation"
            trade["_synthetic_source"] = "corporate_action"
            result.append(trade)
    return result


def _canonical_internal_trade(
    report: ParsedPaidaxReport,
    row: Mapping[str, Any],
    lookup: Mapping[str, Mapping[str, Any]],
    quantity: Decimal,
    trade_type: str,
) -> dict[str, Any]:
    isin = _text(row.get("isin"))
    instrument = lookup.get(isin or "", {})
    symbol = _text(instrument.get("symbol")) or _text(row.get("symbol")) or isin or "UNKNOWN"
    price = abs(_decimal(row.get("price")))
    amount = abs(_decimal(row.get("amount")))
    if amount == 0 and price and quantity:
        amount = abs(quantity * price)
    commission = abs(_decimal(row.get("commission")))
    source_row = row.get("source_row")
    trade_id = f"{report.path.name}:{source_row}"
    if trade_type.startswith("corporate_action:"):
        trade_id = f"CA:{trade_id}"
    return {
        "date_time": row.get("date_time"),
        "trade_id": trade_id,
        "trade_type": trade_type,
        "symbol": symbol,
        "isin": isin,
        "asset_type": _text(instrument.get("type")) or _asset_type(row.get("security_type")),
        "quantity": str(quantity),
        "calculation_quantity": str(quantity),
        "price": str(price),
        "calculation_price": str(price),
        "multiplier": "1",
        "_calculation_multiplier": "1",
        "amount": str(amount),
        "commission": str(commission),
        "amount_with_commission": str(amount + commission),
        "currency": _currency(row.get("currency")) or BASE_CURRENCY,
        "exchange": _exchange(row.get("exchange")),
        "country": _text(instrument.get("country")) or _country(row.get("issuer_country"), isin),
        "source_report": row.get("source_report"),
        "_instrument_identity_key": _instrument_identity_key_from_values(isin=isin, symbol=symbol),
        "_broker_realized_pl": row.get("broker_realized_pl"),
    }


def _build_transfers(
    reports: Sequence[ParsedPaidaxReport], instruments: Sequence[Mapping[str, Any]]
) -> list[dict[str, Any]]:
    lookup = {str(row.get("isin")): row for row in instruments}
    transfers: list[dict[str, Any]] = []
    for report in reports:
        for row in report.cash_transfers:
            amount = _decimal(row.get("amount"))
            transfers.append(
                {
                    "date": row.get("date"),
                    "transfer_type": "cash",
                    "direction": "in" if amount > 0 else "out",
                    "asset_type": "Cash",
                    "symbol": None,
                    "isin": None,
                    "currency": row.get("currency"),
                    "quantity": None,
                    "price": None,
                    "enter_date": None,
                    "amount": str(amount),
                    "broker_comment": row.get("operation"),
                    "counterparty": None,
                    "source_report": row.get("source_report"),
                    "_transfer_id": f"{report.path.name}:{row.get('source_row')}:cash",
                }
            )
        for row in report.security_movements:
            operation = str(row.get("operation") or "").casefold()
            if "безвозмезд" in operation or "подар" in operation:
                continue
            direction = "out" if _is_minus(row.get("direction")) else "in"
            isin = _text(row.get("isin"))
            instrument = lookup.get(isin or "", {})
            quantity = abs(_decimal(row.get("quantity")))
            if quantity == 0:
                continue
            symbol = _text(instrument.get("symbol")) or _text(row.get("symbol")) or isin
            transfers.append(
                {
                    "date": _date_part(row.get("date_time")),
                    "transfer_type": "security",
                    "direction": direction,
                    "asset_type": _text(instrument.get("type")) or _asset_type(row.get("security_type")),
                    "symbol": symbol,
                    "isin": isin,
                    "currency": row.get("currency"),
                    "quantity": str(quantity),
                    "price": row.get("acquisition_price") if direction == "in" else row.get("price"),
                    "enter_date": _date_part(row.get("date_time")) if direction == "in" else None,
                    "amount": None,
                    "broker_comment": row.get("operation"),
                    "counterparty": None,
                    "source_report": row.get("source_report"),
                    "country": _text(instrument.get("country")) or _country(row.get("issuer_country"), isin),
                    "exchange": _exchange(row.get("exchange")),
                    "_raw_quantity": str(quantity if direction == "in" else -quantity),
                    "_transfer_id": f"{report.path.name}:{row.get('source_row')}:security",
                    "_instrument_identity_key": _instrument_identity_key_from_values(isin=isin, symbol=symbol),
                    "_multiplier": "1",
                    "_transfer_cost_basis_status": "broker_reported_cost_basis" if direction == "in" else None,
                    "_fifo_enter_date": row.get("date_time") if direction == "in" else None,
                }
            )
    transfers.extend(_build_position_adjustments(reports, instruments))
    return transfers


def _build_position_adjustments(
    reports: Sequence[ParsedPaidaxReport], instruments: Sequence[Mapping[str, Any]]
) -> list[dict[str, Any]]:
    lookup = {str(row.get("isin")): row for row in instruments}
    latest_end = max((report.period_end for report in reports if report.period_end), default=None)
    result: list[dict[str, Any]] = []
    for report in reports:
        if report.period_end is None or report.period_end != latest_end:
            continue
        expected: dict[str, Decimal] = {
            str(row.get("isin")): _decimal(row.get("starting_quantity"))
            for row in report.positions
            if _text(row.get("isin"))
        }
        row_by_isin: dict[str, Mapping[str, Any]] = {
            str(row.get("isin")): row for row in report.positions if _text(row.get("isin"))
        }
        for row in sorted(
            [*report.trades, *report.security_movements],
            key=lambda item: str(item.get("date_time") or ""),
        ):
            isin = _text(row.get("isin"))
            if not isin:
                continue
            row_by_isin.setdefault(isin, row)
            operation = str(row.get("operation") or "").casefold()
            quantity = abs(_decimal(row.get("quantity")))
            if row.get("source_sheet") == TRADES_SHEET:
                expected[isin] = expected.get(isin, Decimal("0")) + (
                    -quantity if operation == "продажа" else quantity
                )
            elif "безвозмезд" in operation or "подар" in operation:
                if _is_minus(row.get("direction")):
                    expected[isin] = max(expected.get(isin, Decimal("0")), quantity) - quantity
                else:
                    expected[isin] = expected.get(isin, Decimal("0")) + quantity

        for row in report.corporate_actions:
            isin = _text(row.get("isin"))
            action = str(row.get("action_type") or "").casefold()
            if not isin or "денежн" not in action or "компенсац" not in action:
                continue
            row_by_isin.setdefault(isin, row)
            prior = next(
                (
                    candidate
                    for candidate in reversed(report.corporate_actions)
                    if candidate is not row
                    and candidate.get("isin") == isin
                    and str(candidate.get("date") or "") <= str(row.get("date") or "")
                    and (
                        "сплит" in str(candidate.get("action_type") or "").casefold()
                        or "консолидац" in str(candidate.get("action_type") or "").casefold()
                    )
                ),
                None,
            )
            if prior:
                disposed = max(
                    abs(_decimal(prior.get("before_quantity")))
                    - abs(_decimal(prior.get("after_quantity"))),
                    Decimal("0"),
                )
                expected[isin] = max(expected.get(isin, Decimal("0")), disposed) - disposed

        actual = {
            str(row.get("isin")): _decimal(row.get("quantity"))
            for row in report.positions
            if _text(row.get("isin"))
        }
        for isin in sorted(set(expected) | set(actual)):
            difference = actual.get(isin, Decimal("0")) - expected.get(isin, Decimal("0"))
            if difference == 0:
                continue
            source = row_by_isin.get(isin, {})
            instrument = lookup.get(isin, {})
            symbol = _text(instrument.get("symbol")) or _text(source.get("symbol")) or isin
            currency = _currency(source.get("currency")) or BASE_CURRENCY
            direction = "in" if difference > 0 else "out"
            result.append(
                {
                    "date": report.period_end.isoformat(),
                    "transfer_type": "security",
                    "direction": direction,
                    "asset_type": _text(instrument.get("type")) or _asset_type(source.get("security_type")),
                    "symbol": symbol,
                    "isin": isin,
                    "currency": currency,
                    "quantity": str(abs(difference)),
                    "price": "0",
                    "enter_date": report.period_end.isoformat() if direction == "in" else None,
                    "amount": None,
                    "broker_comment": "Paidax ending-position reconciliation adjustment",
                    "counterparty": None,
                    "source_report": str(report.path),
                    "country": _text(instrument.get("country")) or _country(source.get("issuer_country"), isin),
                    "exchange": _exchange(source.get("exchange")),
                    "_raw_quantity": str(difference),
                    "_transfer_id": f"{report.path.name}:position-adjustment:{isin}:{difference}",
                    "_instrument_identity_key": _instrument_identity_key_from_values(isin=isin, symbol=symbol),
                    "_multiplier": "1",
                    "_transfer_cost_basis_status": "summary_position_reconciliation",
                    "_synthetic_reconciliation_adjustment": True,
                }
            )
            report.unprocessed.append(
                _unprocessed(
                    {
                        **source,
                        "source_report": str(report.path),
                        "source_sheet": POSITIONS_SHEET,
                        "date_time": report.period_end.isoformat(),
                        "quantity": str(difference),
                    },
                    "unexplained_ending_position",
                    "Paidax ending position does not reconcile to the report's trades and security movements; a zero-basis inventory adjustment was added.",
                )
            )
    return result


def _build_corporate_actions(
    reports: Sequence[ParsedPaidaxReport], instruments: Sequence[Mapping[str, Any]]
) -> list[dict[str, Any]]:
    lookup = {str(row.get("isin")): row for row in instruments}
    result: list[dict[str, Any]] = []
    for report in reports:
        for row in report.corporate_actions:
            isin = _text(row.get("isin"))
            instrument = lookup.get(isin or "", {})
            action = str(row.get("action_type") or "").casefold()
            action_type = (
                "split_compensation"
                if "денежн" in action and "компенсац" in action
                else "reverse_split"
                if "консолидац" in action or "обратн" in action
                else "split"
                if "сплит" in action
                else "other"
            )
            received = abs(_decimal(row.get("received")))
            result.append(
                {
                    "date": row.get("date"),
                    "symbol": _text(instrument.get("symbol")) or _text(row.get("symbol")) or isin,
                    "isin": isin,
                    "action_type": action_type,
                    "description": row.get("comment") or row.get("action_type"),
                    "quantity": row.get("before_quantity") or row.get("record_quantity"),
                    "proceeds": _money_text(received),
                    "value": _money_text(received),
                    "currency": row.get("currency"),
                    "realized_pl": "0.00",
                    "source_report": row.get("source_report"),
                }
            )
    return result


def _build_dividends(
    reports: Sequence[ParsedPaidaxReport],
    instruments: Sequence[Mapping[str, Any]],
    fx_provider: AnnualFxRateProvider,
    warnings: list[str],
) -> list[dict[str, Any]]:
    lookup = {str(row.get("isin")): row for row in instruments}
    result: list[dict[str, Any]] = []
    for report in reports:
        for row in report.payments:
            isin = _text(row.get("isin"))
            instrument = lookup.get(isin or "", {})
            currency = _currency(row.get("currency")) or BASE_CURRENCY
            gross = _decimal(row.get("gross"))
            withholding = -abs(_decimal(row.get("withholding")))
            net = _decimal(row.get("net"))
            rate = _annual_rate(fx_provider, _year(row.get("date")), currency, warnings)
            country = _text(instrument.get("country")) or _country(row.get("issuer_country"), isin)
            tax = Decimal("0") if country == "KZ" else max(gross, Decimal("0")) * Decimal("0.10")
            result.append(
                {
                    "date": row.get("record_date") or row.get("date"),
                    "pay_date": row.get("date"),
                    "symbol": _text(instrument.get("symbol")) or _text(row.get("symbol")) or isin,
                    "isin": isin,
                    "country": country,
                    "currency": currency,
                    "gross_amount": _money_text(gross),
                    "withholding_tax": _money_text(withholding),
                    "net_amount": _money_text(net),
                    "kzt_rate": str(rate) if rate is not None else None,
                    "gross_amount_kzt": _amount_kzt(gross, rate),
                    "tax": str(tax),
                    "tax_kzt": _amount_kzt(tax, rate),
                    "offshore_flag": False if country == "KZ" else None,
                    "kase_aix_preferential_flag": True if country == "KZ" else None,
                    "source_report": row.get("source_report"),
                }
            )
    return result


def _build_interest(
    reports: Sequence[ParsedPaidaxReport],
    fx_provider: AnnualFxRateProvider,
    warnings: list[str],
) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    for report in reports:
        for row in report.savings:
            gross = _decimal(row.get("income"))
            if gross == 0:
                continue
            currency = _currency(row.get("currency")) or BASE_CURRENCY
            year = report.period_end.year if report.period_end else _year(row.get("date"))
            rate = _annual_rate(fx_provider, year, currency, warnings)
            result.append(
                {
                    "date": row.get("date"),
                    "description": f"Paidax savings income: {row.get('product')}",
                    "financing_kind": "savings_interest",
                    "currency": currency,
                    "gross_amount": _money_text(gross),
                    "withholding_tax": "0.00",
                    "net_amount": _money_text(gross),
                    "kzt_rate": str(rate) if rate is not None else None,
                    "gross_amount_kzt": _amount_kzt(gross, rate),
                    "withholding_tax_kzt": "0.00",
                    "net_amount_kzt": _amount_kzt(gross, rate),
                    "commission": "0.00",
                    "source_report": row.get("source_report"),
                }
            )
    return result


def _build_cash_balances(
    reports: Sequence[ParsedPaidaxReport],
    account_id: str,
    fx_provider: AnnualFxRateProvider,
    warnings: list[str],
) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    for report in reports:
        if report.period_end is None:
            continue
        balances: dict[str, Decimal] = {}
        sources: dict[str, str | None] = {}
        for row in report.cash_balances:
            currency = _currency(row.get("currency")) or BASE_CURRENCY
            balances[currency] = balances.get(currency, Decimal("0")) + _decimal(row.get("ending_cash"))
            sources[currency] = _text(row.get("source_report"))
        for row in report.savings:
            currency = _currency(row.get("currency")) or BASE_CURRENCY
            balances[currency] = balances.get(currency, Decimal("0")) + _decimal(row.get("ending_balance"))
            sources[currency] = _text(row.get("source_report"))
        for currency, amount in sorted(balances.items()):
            rate = _annual_rate(fx_provider, report.period_end.year, currency, warnings)
            result.append(
                {
                    "broker": BROKER_CODE,
                    "account_id": account_id,
                    "year": report.period_end.year,
                    "date": report.period_end.isoformat(),
                    "currency": currency,
                    "ending_cash": _money_text(amount),
                    "ending_cash_kzt": _amount_kzt(amount, rate),
                    "source_report": sources.get(currency),
                }
            )
    return result


def _populate_raw_totals(
    dataset: CanonicalDataset,
    reports: Sequence[ParsedPaidaxReport],
    trades: Sequence[Mapping[str, Any]],
) -> None:
    total_amount = Decimal("0")
    total_commission = Decimal("0")
    turnover_metric = ReconciliationMetric.TRADE_GROSS_AMOUNT_BY_INSTRUMENT.value
    for trade in trades:
        amount = abs(_decimal(trade.get("amount")))
        commission = abs(_decimal(trade.get("commission")))
        total_amount += amount
        total_commission += commission
        key = _dimension_key(
            metric=turnover_metric,
            year=_year(trade.get("date_time")),
            currency=_currency(trade.get("currency")),
            instrument_key=_text(trade.get("isin") or trade.get("symbol")),
        )
        dataset.raw_totals.totals_by_metric_currency[key] = (
            dataset.raw_totals.totals_by_metric_currency.get(key, Decimal("0")) + amount
        )
    dataset.raw_totals.scalar_totals[ReconciliationMetric.TOTAL_TRADES_GROSS_AMOUNT.value] = total_amount
    dataset.raw_totals.scalar_totals[ReconciliationMetric.TOTAL_COMMISSIONS.value] = total_commission
    dividends = dataset.tables.get("Dividends", [])
    dividend_gross = sum((_decimal(row.get("gross_amount")) for row in dividends), Decimal("0"))
    dividend_tax = sum((_decimal(row.get("withholding_tax")) for row in dividends), Decimal("0"))
    dataset.raw_totals.scalar_totals[ReconciliationMetric.TOTAL_DIVIDENDS_GROSS.value] = dividend_gross
    dataset.raw_totals.scalar_totals[ReconciliationMetric.TOTAL_DIVIDENDS_TAX.value] = dividend_tax
    dataset.raw_totals.scalar_totals[ReconciliationMetric.TOTAL_DIVIDENDS_NET.value] = dividend_gross + dividend_tax
    dataset.raw_totals.scalar_totals[ReconciliationMetric.TOTAL_INTEREST.value] = sum(
        (_decimal(row.get("gross_amount")) for row in dataset.tables.get("Interest", [])), Decimal("0")
    )
    transfer_total = sum(
        (_decimal(row.get("amount")) for row in dataset.tables.get("Transfers", [])), Decimal("0")
    )
    dataset.raw_totals.scalar_totals[
        ReconciliationMetric.TOTAL_DEPOSITS_WITHDRAWALS_TRANSFERS.value
    ] = transfer_total
    for row in dataset.tables.get("Transfers", []):
        if row.get("amount") in (None, ""):
            continue
        key = _dimension_key(
            metric=ReconciliationMetric.TOTAL_DEPOSITS_WITHDRAWALS_TRANSFERS.value,
            currency=_currency(row.get("currency")),
        )
        dataset.raw_totals.totals_by_metric_currency[key] = (
            dataset.raw_totals.totals_by_metric_currency.get(key, Decimal("0"))
            + _decimal(row.get("amount"))
        )
    for row in dataset.tables.get("CashBalances", []):
        key = _dimension_key(year=int(row["year"]), currency=_currency(row.get("currency")))
        dataset.raw_totals.cash_by_currency[key] = (
            dataset.raw_totals.cash_by_currency.get(key, Decimal("0"))
            + _decimal(row.get("ending_cash"))
        )
    for report in reports:
        if report.period_end is None:
            continue
        for row in report.positions:
            isin = _text(row.get("isin"))
            if not isin:
                continue
            key = _dimension_key(year=report.period_end.year, instrument_key=isin)
            dataset.raw_totals.positions_by_key[key] = (
                dataset.raw_totals.positions_by_key.get(key, Decimal("0"))
                + _decimal(row.get("quantity"))
            )


def _security_row(
    row: Mapping[str, Any], path: Path, source_row: int, source_sheet: str
) -> dict[str, Any]:
    return {
        "symbol": _text(row.get("Тикер")),
        "name": _text(row.get("Наименование")),
        "isin": _text(row.get("ISIN")),
        "issuer_country": _text(row.get("Код страны регистрации эмитента")),
        "exchange": _text(row.get("Наименование рынка")),
        "security_type": _text(row.get("Тип актива") or row.get("Тип ценной бумаги")),
        "currency": _currency(row.get("Валюта")),
        "source_report": str(path),
        "source_row": source_row,
        "source_sheet": source_sheet,
    }


def _is_crypto(row: Mapping[str, Any]) -> bool:
    asset_type = _clean_text(row.get("security_type")).casefold()
    if any(token in asset_type for token in ("крипт", "crypto", "digital asset", "цифровой актив")):
        return True
    if _text(row.get("isin")):
        return False
    symbol = (_text(row.get("symbol")) or "").upper().replace(" ", "")
    base_symbol = re.split(r"[/_.:-]", symbol, maxsplit=1)[0]
    if base_symbol in _DIRECT_CRYPTO_SYMBOLS:
        return True
    return any(
        symbol.startswith(f"{crypto}{quote}")
        for crypto in _DIRECT_CRYPTO_SYMBOLS
        for quote in ("USD", "USDT", "USDC", "KZT", "EUR")
    )


def _crypto_unprocessed(row: Mapping[str, Any]) -> dict[str, Any]:
    return _unprocessed(
        row,
        "unsupported_crypto_asset",
        "Direct cryptocurrency operations require a separate tax calculation and were excluded from securities processing.",
    )


def _unprocessed(
    row: Mapping[str, Any], reason: str, details: str
) -> dict[str, Any]:
    source_report = _text(row.get("source_report"))
    return {
        "severity": "warning",
        "reason": reason,
        "details": details,
        "source_sheet": row.get("source_sheet") or TRADES_SHEET,
        "source_report": source_report,
        "trade_id": f"{Path(source_report).name}:{row.get('source_row')}" if source_report else None,
        "date_time": row.get("date_time") or row.get("date"),
        "symbol": row.get("symbol"),
        "isin": row.get("isin"),
        "asset_type": row.get("security_type"),
        "currency": row.get("currency"),
        "quantity": row.get("quantity"),
        "price": row.get("price"),
        "amount": row.get("amount") or row.get("gross") or row.get("received"),
        "commission": row.get("commission"),
    }


def _records(
    worksheet: Any, required_headers: Sequence[str]
) -> list[tuple[int, dict[str, Any]]]:
    header_row = _find_header_row(worksheet, required_headers)
    if header_row is None:
        return []
    headers = _headers(worksheet, header_row)
    result: list[tuple[int, dict[str, Any]]] = []
    for row_index, values in enumerate(
        worksheet.iter_rows(min_row=header_row + 1, values_only=True),
        start=header_row + 1,
    ):
        if not any(value not in (None, "") for value in values):
            continue
        result.append((row_index, _row_mapping(headers, values)))
    return result


def _find_header_row(worksheet: Any, required_headers: Sequence[str]) -> int | None:
    required = {_normalize_header(header) for header in required_headers}
    for row_index in range(1, min(30, worksheet.max_row) + 1):
        actual = {
            _normalize_header(worksheet.cell(row_index, column).value)
            for column in range(1, worksheet.max_column + 1)
        }
        if required.issubset(actual):
            return row_index
    return None


def _headers(worksheet: Any, row_index: int) -> dict[str, int]:
    return {
        _normalize_header(worksheet.cell(row_index, column).value): column - 1
        for column in range(1, worksheet.max_column + 1)
        if _normalize_header(worksheet.cell(row_index, column).value)
    }


def _row_mapping(headers: Mapping[str, int], values: Sequence[Any]) -> dict[str, Any]:
    return {header: values[index] if index < len(values) else None for header, index in headers.items()}


def _normalize_header(value: Any) -> str:
    return re.sub(r"\s+", " ", _clean_text(value)).strip()


def _clean_text(value: Any) -> str:
    return str(value or "").replace("\u00a0", " ").strip()


def _text(value: Any) -> str | None:
    text = _clean_text(value)
    return None if not text or text in {"—", "-"} else text


def _decimal(value: Any) -> Decimal:
    if value in (None, "", "—", "-"):
        return Decimal("0")
    text = str(value).replace("\u00a0", "").replace(" ", "").replace(",", ".").strip()
    try:
        return Decimal(text or "0")
    except InvalidOperation:
        return Decimal("0")


def _optional_decimal_text(value: Any) -> str | None:
    if value in (None, "", "—", "-"):
        return None
    return str(_decimal(value))


def _parse_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _clean_text(value)
    for fmt in ("%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text[:10], fmt).date()
        except ValueError:
            pass
    return None


def _datetime_text(value: Any) -> str | None:
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    parsed = _parse_date(value)
    return datetime.combine(parsed, datetime.min.time()).isoformat(sep=" ") if parsed else None


def _parse_datetime(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    text = _clean_text(value)
    if text:
        try:
            return datetime.fromisoformat(text)
        except ValueError:
            pass
    parsed = _parse_date(value)
    return datetime.combine(parsed, datetime.min.time()) if parsed else None


def _date_text(value: Any) -> str | None:
    parsed = _parse_date(value)
    return parsed.isoformat() if parsed else None


def _date_part(value: Any) -> str | None:
    text = _clean_text(value)
    return text[:10] if len(text) >= 10 else None


def _currency(value: Any) -> str | None:
    text = _text(value)
    return text.upper() if text else None


def _country(value: Any, isin: str | None) -> str | None:
    raw = (_text(value) or "").upper()
    aliases = {"KAZ": "KZ", "USA": "US", "CAN": "CA"}
    if raw:
        return aliases.get(raw, raw)
    return isin[:2].upper() if isin and len(isin) >= 2 else None


def _exchange(value: Any) -> str | None:
    exchange = _text(value)
    return exchange.upper() if exchange else None


def _asset_type(value: Any) -> str:
    normalized = _clean_text(value).casefold()
    if "облигац" in normalized or "bond" in normalized:
        return "Bonds"
    return "Stocks"


def _is_minus(value: Any) -> bool:
    text = _clean_text(value)
    return text.startswith(("-", "−", "–"))


def _max_report_year(reports: Sequence[ParsedPaidaxReport]) -> int | None:
    return max((report.period_end.year for report in reports if report.period_end), default=None)


def _year(value: Any) -> int | None:
    parsed = _parse_date(value)
    return parsed.year if parsed else None


def _annual_rate(
    fx_provider: AnnualFxRateProvider,
    year: int | None,
    currency: str | None,
    warnings: list[str],
) -> Decimal | None:
    if year is None or not currency:
        return None
    rate = fx_provider.rate(year, currency)
    if rate is None:
        warning = f"Missing annual NBK FX rate for {currency}/{year}; KZT fields left empty."
        if warning not in warnings:
            warnings.append(warning)
    return rate


def _money_text(value: Decimal) -> str:
    return str(value.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def _amount_kzt(amount: Decimal, rate: Decimal | None) -> str | None:
    return str(amount * rate) if rate is not None else None


def _later_date(left: date | None, right: date | None) -> date | None:
    if left is None:
        return right
    if right is None:
        return left
    return max(left, right)


def _dimension_key(
    *,
    metric: str | None = None,
    year: int | None = None,
    currency: str | None = None,
    instrument_key: str | None = None,
) -> str:
    return "|".join(
        "" if value is None else str(value)
        for value in (metric, year, currency, instrument_key)
    )
