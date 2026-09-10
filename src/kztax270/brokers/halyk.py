"""Parser for Halyk Finance tax-report workbooks."""

from __future__ import annotations

import re
from collections.abc import Mapping, Sequence
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from kztax270.canonical.schema import AccountMetadata, CanonicalDataset
from kztax270.canonical.trade_enrichment import enrich_trades_before_calculations
from kztax270.diagnostics import instrument_parsed_reports
from kztax270.reconciliation.models import ReconciliationMetric
from kztax270.reference.fx import AnnualFxRateProvider

from .base import BrokerReport, ParseResult
from .discovery import DiscoveryRule, discover_raw_reports
from .ib import (
    _build_broker_trade_realized_pl,
    _build_fifo_and_positions,
    _build_unprocessed_rows,
    _build_years_results,
    _canonical_trade_rows,
    _canonical_transfer_rows,
    _instrument_identity_key_from_values,
    _instrument_symbol_history,
    _sort_trades_by_datetime,
)

BROKER_CODE = "halyk"
RAW_FOLDER = "halyk"
BASE_CURRENCY = "KZT"

MOVEMENT_SHEET = "Движение ФИ"
INCOME_SHEET = "Дивиденды"
REPO_SHEET = "РЕПО"
POSITIONS_SHEET = "ЦБ"

ACCOUNT_LABEL = "Номер лицевого счета"
AS_OF_LABEL = "По состоянию на"
PERIOD_LABEL = "Период за который предоставляется отчет"
CASH_BALANCE_MARKER = "Остатки денежных средств"

_DATE_PATTERN = re.compile(r"\d{2}\.\d{2}\.\d{4}")


@dataclass(slots=True)
class ParsedHalykReport:
    path: Path
    account_id: str | None = None
    period_start: date | None = None
    period_end: date | None = None
    trades: list[dict[str, Any]] = field(default_factory=list)
    security_transfers: list[dict[str, Any]] = field(default_factory=list)
    income: list[dict[str, Any]] = field(default_factory=list)
    positions: list[dict[str, Any]] = field(default_factory=list)
    cash_balances: list[dict[str, Any]] = field(default_factory=list)
    repo: list[dict[str, Any]] = field(default_factory=list)
    unprocessed: list[dict[str, Any]] = field(default_factory=list)


class HalykParser:
    broker_code = BROKER_CODE

    def __init__(self, fx_provider: AnnualFxRateProvider | None = None) -> None:
        self.fx_provider = fx_provider or AnnualFxRateProvider({})

    def discover_reports(self, raw_root: Path, account_id: str) -> list[BrokerReport]:
        return discover_raw_reports(
            raw_root,
            DiscoveryRule(
                broker=RAW_FOLDER,
                account_id=account_id,
                extensions=frozenset({".xlsx"}),
            ),
        )

    def parse_reports(self, reports: Sequence[BrokerReport], account_id: str) -> ParseResult:
        parsed_reports = [parse_halyk_xlsx(report.path) for report in reports]
        instrument_parsed_reports(self.broker_code, parsed_reports)
        dataset = build_canonical_dataset(parsed_reports, account_id, self.fx_provider)
        dataset.raw_totals.source_reports = [str(report.path) for report in reports]
        return ParseResult(
            broker=self.broker_code,
            account_id=account_id,
            reports=reports,
            dataset=dataset,
            raw_totals=dataset.raw_totals,
        )


def parse_halyk_xlsx(path: Path) -> ParsedHalykReport:
    workbook = load_workbook(path, read_only=True, data_only=True)
    parsed = ParsedHalykReport(path=path)
    try:
        for worksheet in workbook.worksheets:
            _read_metadata(worksheet, parsed)
            headers = _header_map(worksheet)
            if _is_movement_sheet(worksheet.title, headers):
                _parse_movements(worksheet, headers, parsed)
            elif _is_income_sheet(worksheet.title, headers):
                _parse_income(worksheet, headers, parsed)
            elif _is_positions_sheet(worksheet.title, headers):
                _parse_positions(worksheet, headers, parsed)
            elif _is_repo_sheet(worksheet.title, headers):
                _parse_repo(worksheet, headers, parsed)
    finally:
        workbook.close()
    return parsed


def build_canonical_dataset(
    reports: Sequence[ParsedHalykReport],
    account_id: str,
    fx_provider: AnnualFxRateProvider,
) -> CanonicalDataset:
    dataset = CanonicalDataset(
        metadata=AccountMetadata(broker=BROKER_CODE, account_id=account_id, base_currency=BASE_CURRENCY)
    )
    for report in reports:
        if report.account_id and not _same_account_id(report.account_id, account_id):
            dataset.warnings.append(
                f"Halyk report {report.path} belongs to account {report.account_id}, expected {account_id}."
            )

    instruments = _build_instruments(reports, account_id)
    dataset.tables["Instruments"] = instruments
    internal_trades = _sort_trades_by_datetime(
        _build_trades(reports, instruments, fx_provider, dataset.warnings)
    )
    enrich_trades_before_calculations(dataset, internal_trades, fx_provider)
    transfers = _build_security_transfers(reports, instruments)
    dataset.tables["Trades"] = _canonical_trade_rows(internal_trades)
    dataset.tables["_BrokerTradeRealizedPL"] = _build_broker_trade_realized_pl(internal_trades)

    fifo_rows, positions, transfer_rows = _build_fifo_and_positions(
        internal_trades,
        transfers=transfers,
        initial_lots=[],
        max_year=_max_report_year(reports),
        fx_provider=fx_provider,
        warnings=dataset.warnings,
        symbol_history=_instrument_symbol_history(instruments),
        join_security_currencies=True,
    )
    dataset.tables["Fifo"] = fifo_rows
    dataset.tables["Positions"] = positions
    dataset.tables["Transfers"] = _canonical_transfer_rows(transfer_rows)
    dataset.tables["CorporateActions"] = _build_corporate_actions(reports, instruments)
    dataset.tables["Dividends"] = _build_dividends(
        reports, instruments, fx_provider, dataset.warnings
    )
    dataset.tables["Interest"] = _build_interest(reports, fx_provider, dataset.warnings)
    dataset.tables["Coupons"] = _build_coupons(
        reports, instruments, fx_provider, dataset.warnings
    )
    dataset.tables["CashBalances"] = _build_cash_balances(
        reports, account_id, fx_provider, dataset.warnings
    )
    dataset.tables["Unprocessed"] = [
        *_build_unprocessed_rows(dataset.tables["Trades"], fifo_rows),
        *(row for report in reports for row in report.unprocessed),
    ]
    dataset.tables["Years_Results"] = _build_years_results(dataset)
    _populate_raw_totals(dataset, reports, internal_trades)
    return dataset


def _read_metadata(worksheet: Any, parsed: ParsedHalykReport) -> None:
    for row in worksheet.iter_rows(min_row=1, max_row=min(8, worksheet.max_row), values_only=True):
        values = [_clean_text(value) for value in row]
        for index, value in enumerate(values):
            normalized = value.casefold().rstrip(":")
            following = next((item for item in values[index + 1 :] if item), "")
            if normalized == ACCOUNT_LABEL.casefold():
                parsed.account_id = following or parsed.account_id
            elif normalized == AS_OF_LABEL.casefold() and following:
                report_date = _parse_date(following)
                if report_date is not None:
                    parsed.period_end = _later_date(parsed.period_end, report_date)
            elif normalized.startswith(PERIOD_LABEL.casefold()) and following:
                period_start, period_end = _parse_period(following)
                parsed.period_start = _earlier_date(parsed.period_start, period_start)
                parsed.period_end = _later_date(parsed.period_end, period_end)


def _header_map(worksheet: Any) -> dict[str, int]:
    for row_index in range(1, min(20, worksheet.max_row) + 1):
        values = [_clean_text(worksheet.cell(row_index, column).value) for column in range(1, worksheet.max_column + 1)]
        if "№" in values and len([value for value in values if value]) >= 3:
            return {value: index for index, value in enumerate(values) if value}
    return {}


def _is_movement_sheet(title: str, headers: Mapping[str, int]) -> bool:
    return title.strip() == MOVEMENT_SHEET or "Направление операции" in headers


def _is_income_sheet(title: str, headers: Mapping[str, int]) -> bool:
    return title.strip() == INCOME_SHEET or (
        "Тип операции" in headers and "Сумма в валюте" in headers
    )


def _is_positions_sheet(title: str, headers: Mapping[str, int]) -> bool:
    return title.strip() == POSITIONS_SHEET or (
        "Количество ЦБ" in headers and "Адрес регистрации" in headers
    )


def _is_repo_sheet(title: str, headers: Mapping[str, int]) -> bool:
    return title.strip() == REPO_SHEET or "Объем открытия в валюте" in headers


def _data_rows(worksheet: Any, headers: Mapping[str, int]) -> Sequence[tuple[int, list[Any]]]:
    if not headers:
        return []
    header_row = next(
        (
            row
            for row in range(1, min(20, worksheet.max_row) + 1)
            if _clean_text(worksheet.cell(row, 1).value) == "№"
        ),
        8,
    )
    return [
        (row_index, [worksheet.cell(row_index, column).value for column in range(1, worksheet.max_column + 1)])
        for row_index in range(header_row + 1, worksheet.max_row + 1)
        if _is_data_number(worksheet.cell(row_index, 1).value)
    ]


def _parse_movements(worksheet: Any, headers: Mapping[str, int], parsed: ParsedHalykReport) -> None:
    for source_row, values in _data_rows(worksheet, headers):
        name = _column(values, headers, "Наименование (тип)")
        currency = _currency(_column(values, headers, "Код валюты"))
        if _clean_text(name).casefold().startswith(CASH_BALANCE_MARKER.casefold()):
            parsed.cash_balances.append(
                {
                    "currency": currency or BASE_CURRENCY,
                    "ending_cash": str(_decimal(_column(values, headers, "Цена приобретения одной ЦБ"))),
                    "source_report": str(parsed.path),
                    "source_row": source_row,
                }
            )
            continue

        isin = _text(_column(values, headers, "ISIN ЦБ"))
        operation = _text(_column(values, headers, "Направление операции"))
        trade_date = _parse_date(_column(values, headers, "Дата торговой операции"))
        if not isin or not operation or trade_date is None:
            continue
        isin_columns = [index for header, index in headers.items() if header.startswith("ISIN")]
        converted_isin = _text(values[isin_columns[1]]) if len(isin_columns) > 1 else None
        operation_kind = (
            "transfer_in"
            if _is_self_security_transfer_in(operation, isin, converted_isin)
            else _movement_kind(operation)
        )
        quantity = _decimal(_column(values, headers, "Количество ЦБ"))
        if operation_kind in {"sale", "redemption"}:
            quantity = -abs(quantity)
        elif operation_kind == "purchase":
            quantity = abs(quantity)
        row = {
            "date_time": datetime.combine(trade_date, datetime.min.time()).isoformat(sep=" "),
            "name": _clean_text(name),
            "symbol": _symbol_from_name(name, isin),
            "security_type": _column(values, headers, "Тип ЦБ"),
            "isin": isin,
            "operation": operation,
            "operation_kind": operation_kind,
            "quantity": str(quantity),
            "conversion_coefficient": str(_decimal(values[14])) if len(values) > 14 else "0",
            "converted_isin": converted_isin,
            "transfer_from": _text(values[6]) if len(values) > 6 else None,
            "price": str(_decimal(_column(values, headers, "Цена приобретения одной ЦБ"))),
            "currency": currency or BASE_CURRENCY,
            "exchange": _movement_exchange(values, headers),
            "issuer_country": _text(_column(values, headers, "Код страны регистрации эмитента")),
            "commission_kzt": str(abs(_decimal(_column(values, headers, "Комиссии по сделке")))),
            "order_id": _text(_column(values, headers, "Номер заказа")),
            "source_report": str(parsed.path),
            "source_row": source_row,
        }
        if operation_kind is None:
            parsed.unprocessed.append(_unprocessed(row, "unsupported_operation", f"Unsupported Halyk security operation: {operation}"))
        elif operation_kind == "transfer_in":
            parsed.security_transfers.append(row)
        else:
            parsed.trades.append(row)


def _parse_income(worksheet: Any, headers: Mapping[str, int], parsed: ParsedHalykReport) -> None:
    for source_row, values in _data_rows(worksheet, headers):
        paid_at = _parse_date(_column(values, headers, "Дата транзакции"))
        if paid_at is None:
            continue
        operation = _text(_column(values, headers, "Тип операции")) or ""
        isin = _text(_column(values, headers, "ISIN ЦБ"))
        name = _column(values, headers, "Наименование (тип)")
        row = {
            "date": paid_at.isoformat(),
            "name": _clean_text(name),
            "symbol": _symbol_from_name(name, isin),
            "isin": isin,
            "quantity": str(_decimal(_column(values, headers, "Количество ЦБ"))),
            "currency": _currency(_column(values, headers, "Код валюты")) or BASE_CURRENCY,
            "amount": str(_decimal(_column(values, headers, "Сумма в валюте"))),
            "operation": operation,
            "security_type": _column(values, headers, "Тип ЦБ"),
            "exchange": _text(_column(values, headers, "Рынок ЦБ")),
            "issuer_country": _text(_column(values, headers, "Код страны регистрации эмитента")),
            "source_report": str(parsed.path),
            "source_row": source_row,
        }
        operation_key = operation.casefold()
        if "купон" in operation_key or "дивиденд" in operation_key or "вознагражден" in operation_key:
            parsed.income.append(row)
        else:
            parsed.unprocessed.append(_unprocessed(row, "unsupported_income", f"Unsupported Halyk income operation: {operation}"))


def _parse_positions(worksheet: Any, headers: Mapping[str, int], parsed: ParsedHalykReport) -> None:
    for source_row, values in _data_rows(worksheet, headers):
        isin = _text(_column(values, headers, "ISIN ЦБ"))
        if not isin:
            continue
        name = _column(values, headers, "Наименование")
        parsed.positions.append(
            {
                "name": _clean_text(name),
                "symbol": _symbol_from_name(name, isin),
                "isin": isin,
                "security_type": _column(values, headers, "Тип ЦБ"),
                "issuer_country": _text(_column(values, headers, "Код страны регистрации эмитента")),
                "quantity": str(_decimal(_column(values, headers, "Количество ЦБ"))),
                "currency": _currency(_column(values, headers, "Код валюты")) or BASE_CURRENCY,
                "source_report": str(parsed.path),
                "source_row": source_row,
            }
        )


def _parse_repo(worksheet: Any, headers: Mapping[str, int], parsed: ParsedHalykReport) -> None:
    for source_row, values in _data_rows(worksheet, headers):
        row = {
            "date_time": _date_text(_column(values, headers, "Дата транзакции")),
            "symbol": _text(_column(values, headers, "Наименование (тип)")),
            "currency": _currency(_column(values, headers, "Код валюты")) or BASE_CURRENCY,
            "quantity": str(_decimal(_column(values, headers, "Количество ЦБ"))),
            "amount": str(
                abs(_decimal(_column(values, headers, "Объем открытия в валюте")))
                + abs(_decimal(_column(values, headers, "Объем закрытия в валюте")))
            ),
            "commission": str(abs(_decimal(_column(values, headers, "Комиссии по сделке")))),
            "source_report": str(parsed.path),
            "source_row": source_row,
        }
        parsed.repo.append(row)
        parsed.unprocessed.append(
            _unprocessed(row, "unsupported_repo", "Halyk REPO row was preserved but is not included in tax trades.")
        )


def _build_instruments(
    reports: Sequence[ParsedHalykReport], account_id: str
) -> list[dict[str, Any]]:
    sources: dict[str, Mapping[str, Any]] = {}
    latest_dates: dict[str, date | None] = {}
    for report in reports:
        for row in [*report.trades, *report.income, *report.positions]:
            isin = _text(row.get("isin"))
            if not isin:
                continue
            existing = sources.get(isin)
            if existing is None or (not _text(existing.get("exchange")) and _text(row.get("exchange"))):
                sources[isin] = row
            latest_dates[isin] = _later_date(latest_dates.get(isin), report.period_end)

    instruments: list[dict[str, Any]] = []
    for isin in sorted(sources):
        source = sources[isin]
        country = _country(source.get("issuer_country"), isin)
        symbol = _text(source.get("symbol")) or isin
        instruments.append(
            {
                "symbol": symbol,
                "description": _text(source.get("name")) or isin,
                "conid": None,
                "security_id": isin,
                "underlying": None,
                "listing_exchange": _normalize_exchange(source.get("exchange")),
                "multiplier": "1",
                "type": _asset_type(source.get("security_type")),
                "code": None,
                "year": None,
                "expiry": None,
                "delivery_month": None,
                "strike": None,
                "issuer": _text(source.get("name")),
                "maturity": None,
                "cusip": None,
                "country": country,
                "isin": isin,
                "figi": None,
                "issuer_country": country,
                "offshore_flag": False if country == "KZ" else None,
                "issuer_outside_kz_flag": False if country == "KZ" else (True if country else None),
                "preferential_tax_flag": None,
                "source_broker": BROKER_CODE,
                "source_account": account_id,
                "source_report": source.get("source_report"),
                "as_of_date": latest_dates[isin].isoformat() if latest_dates.get(isin) else None,
            }
        )
    return instruments


def _build_trades(
    reports: Sequence[ParsedHalykReport],
    instruments: Sequence[Mapping[str, Any]],
    fx_provider: AnnualFxRateProvider,
    warnings: list[str],
) -> list[dict[str, Any]]:
    lookup = {str(row.get("isin")): row for row in instruments}
    trades: list[dict[str, Any]] = []
    for report in reports:
        for row in report.trades:
            isin = _text(row.get("isin"))
            instrument = lookup.get(isin or "", {})
            quantity = _decimal(row.get("quantity"))
            price = _decimal(row.get("price"))
            amount = abs(quantity * price)
            currency = _currency(row.get("currency")) or BASE_CURRENCY
            commission = _commission_in_trade_currency(
                _decimal(row.get("commission_kzt")),
                _year(row.get("date_time")),
                currency,
                fx_provider,
                warnings,
            )
            symbol = _text(instrument.get("symbol")) or isin
            country = _text(instrument.get("country")) or _country(row.get("issuer_country"), isin)
            is_redemption = row.get("operation_kind") == "redemption"
            trade_id = f"{report.path.name}:{row.get('source_row')}"
            if is_redemption:
                trade_id = f"CA:{trade_id}"
            trade = {
                "date_time": row.get("date_time"),
                "trade_id": trade_id,
                "trade_type": "redemption" if is_redemption else "trade",
                "symbol": symbol,
                "isin": isin,
                "asset_type": _text(instrument.get("type")) or _asset_type(row.get("security_type")),
                "quantity": str(quantity),
                "calculation_quantity": str(quantity),
                "price": str(price),
                "calculation_price": str(price),
                "multiplier": "1",
                "_calculation_multiplier": "1",
                "amount": str(amount),
                "commission": str(commission),
                "amount_with_commission": str(amount + commission),
                "currency": currency,
                "exchange": _normalize_exchange(row.get("exchange")),
                "country": country,
                "source_report": row.get("source_report"),
                "_instrument_identity_key": _instrument_identity_key_from_values(isin=isin, symbol=symbol),
                "_broker_realized_pl": "0",
            }
            if is_redemption:
                trade["_corporate_action_type"] = "maturity"
                trade["corporate_action_type"] = "maturity"
            trades.append(trade)
    return trades


def _build_security_transfers(
    reports: Sequence[ParsedHalykReport],
    instruments: Sequence[Mapping[str, Any]],
) -> list[dict[str, Any]]:
    lookup = {str(row.get("isin")): row for row in instruments}
    transfers: list[dict[str, Any]] = []
    for report in reports:
        for row in report.security_transfers:
            isin = _text(row.get("isin"))
            instrument = lookup.get(isin or "", {})
            quantity = abs(_decimal(row.get("quantity")))
            if quantity == 0:
                continue
            price = _decimal(row.get("conversion_coefficient"))
            symbol = _text(instrument.get("symbol")) or isin
            transfers.append(
                {
                    "date": _date_part(row.get("date_time")),
                    "transfer_type": "security",
                    "direction": "in",
                    "asset_type": _text(instrument.get("type")) or _asset_type(row.get("security_type")),
                    "symbol": symbol,
                    "isin": isin,
                    "currency": _currency(row.get("currency")) or BASE_CURRENCY,
                    "quantity": str(quantity),
                    "price": str(price),
                    "enter_date": _date_part(row.get("date_time")),
                    "amount": None,
                    "broker_comment": (
                        f"{row.get('operation')}; self-ISIN transfer-in; "
                        f"cost basis from conversion coefficient: {price}"
                    ),
                    "counterparty": _text(row.get("transfer_from")),
                    "source_report": row.get("source_report"),
                    "country": _text(instrument.get("country")) or _country(row.get("issuer_country"), isin),
                    "exchange": _normalize_exchange(row.get("exchange")),
                    "_raw_quantity": str(quantity),
                    "_transfer_id": f"{report.path.name}:{row.get('source_row')}:transfer-in",
                    "_instrument_identity_key": _instrument_identity_key_from_values(isin=isin, symbol=symbol),
                    "_multiplier": "1",
                    "_transfer_cost_basis_status": "broker_reported_cost_basis",
                    "_fifo_enter_date": row.get("date_time"),
                }
            )
    return transfers


def _build_corporate_actions(
    reports: Sequence[ParsedHalykReport], instruments: Sequence[Mapping[str, Any]]
) -> list[dict[str, Any]]:
    lookup = {str(row.get("isin")): row for row in instruments}
    actions: list[dict[str, Any]] = []
    for report in reports:
        for row in report.trades:
            if row.get("operation_kind") != "redemption":
                continue
            isin = _text(row.get("isin"))
            instrument = lookup.get(isin or "", {})
            quantity = _decimal(row.get("quantity"))
            proceeds = abs(quantity * _decimal(row.get("price")))
            actions.append(
                {
                    "date": _date_part(row.get("date_time")),
                    "symbol": _text(instrument.get("symbol")) or isin,
                    "isin": isin,
                    "action_type": "maturity",
                    "description": row.get("operation"),
                    "quantity": str(quantity),
                    "proceeds": _money_text(proceeds),
                    "value": _money_text(proceeds),
                    "currency": row.get("currency"),
                    "realized_pl": "0.00",
                    "source_report": row.get("source_report"),
                }
            )
    return actions


def _build_dividends(
    reports: Sequence[ParsedHalykReport],
    instruments: Sequence[Mapping[str, Any]],
    fx_provider: AnnualFxRateProvider,
    warnings: list[str],
) -> list[dict[str, Any]]:
    lookup = {str(row.get("isin")): row for row in instruments}
    result: list[dict[str, Any]] = []
    for report in reports:
        for row in report.income:
            if "дивиденд" not in str(row.get("operation") or "").casefold():
                continue
            isin = _text(row.get("isin"))
            instrument = lookup.get(isin or "", {})
            currency = _currency(row.get("currency")) or BASE_CURRENCY
            gross = _decimal(row.get("amount"))
            rate = _annual_rate(fx_provider, _year(row.get("date")), currency, warnings)
            country = _text(instrument.get("country")) or _country(row.get("issuer_country"), isin)
            tax = Decimal("0") if country == "KZ" else max(gross, Decimal("0")) * Decimal("0.10")
            result.append(
                {
                    "date": row.get("date"),
                    "pay_date": row.get("date"),
                    "symbol": _text(instrument.get("symbol")) or isin,
                    "isin": isin,
                    "country": country,
                    "currency": currency,
                    "gross_amount": _money_text(gross),
                    "withholding_tax": "0.00",
                    "net_amount": _money_text(gross),
                    "kzt_rate": str(rate) if rate is not None else None,
                    "gross_amount_kzt": _amount_kzt(gross, rate),
                    "tax": str(tax),
                    "tax_kzt": _amount_kzt(tax, rate),
                    "offshore_flag": False if country == "KZ" else None,
                    "kase_aix_preferential_flag": True if country == "KZ" else None,
                    "source_report": row.get("source_report"),
                }
            )
    return result


def _build_coupons(
    reports: Sequence[ParsedHalykReport],
    instruments: Sequence[Mapping[str, Any]],
    fx_provider: AnnualFxRateProvider,
    warnings: list[str],
) -> list[dict[str, Any]]:
    lookup = {str(row.get("isin")): row for row in instruments}
    result: list[dict[str, Any]] = []
    for report in reports:
        for row in report.income:
            if "купон" not in str(row.get("operation") or "").casefold():
                continue
            isin = _text(row.get("isin"))
            instrument = lookup.get(isin or "", {})
            currency = _currency(row.get("currency")) or BASE_CURRENCY
            gross = _decimal(row.get("amount"))
            rate = _annual_rate(fx_provider, _year(row.get("date")), currency, warnings)
            country = _text(instrument.get("country")) or _country(row.get("issuer_country"), isin)
            result.append(
                {
                    "date": row.get("date"),
                    "symbol": _text(instrument.get("symbol")) or isin,
                    "isin": isin,
                    "country": country,
                    "currency": currency,
                    "gross_amount": _money_text(gross),
                    "withholding_tax": "0.00",
                    "net_amount": _money_text(gross),
                    "kzt_rate": str(rate) if rate is not None else None,
                    "gross_amount_kzt": _amount_kzt(gross, rate),
                    "withholding_tax_kzt": "0.00",
                    "net_amount_kzt": _amount_kzt(gross, rate),
                    "is_revert": gross < 0,
                    "offshore_flag": False if country == "KZ" else None,
                    "source_report": row.get("source_report"),
                }
            )
    return result


def _build_interest(
    reports: Sequence[ParsedHalykReport],
    fx_provider: AnnualFxRateProvider,
    warnings: list[str],
) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    for report in reports:
        for row in report.income:
            operation_key = str(row.get("operation") or "").casefold()
            if "вознагражден" not in operation_key or "купон" in operation_key or "дивиденд" in operation_key:
                continue
            currency = _currency(row.get("currency")) or BASE_CURRENCY
            gross = _decimal(row.get("amount"))
            rate = _annual_rate(fx_provider, _year(row.get("date")), currency, warnings)
            result.append(
                {
                    "date": row.get("date"),
                    "description": row.get("operation"),
                    "financing_kind": "cash_interest",
                    "currency": currency,
                    "gross_amount": _money_text(gross),
                    "withholding_tax": "0.00",
                    "net_amount": _money_text(gross),
                    "kzt_rate": str(rate) if rate is not None else None,
                    "gross_amount_kzt": _amount_kzt(gross, rate),
                    "withholding_tax_kzt": "0.00",
                    "net_amount_kzt": _amount_kzt(gross, rate),
                    "commission": "0.00",
                    "source_report": row.get("source_report"),
                }
            )
    return result


def _build_cash_balances(
    reports: Sequence[ParsedHalykReport],
    account_id: str,
    fx_provider: AnnualFxRateProvider,
    warnings: list[str],
) -> list[dict[str, Any]]:
    balances: list[dict[str, Any]] = []
    for report in reports:
        if report.period_end is None:
            continue
        for row in report.cash_balances:
            currency = _currency(row.get("currency")) or BASE_CURRENCY
            ending_cash = _decimal(row.get("ending_cash"))
            rate = _annual_rate(fx_provider, report.period_end.year, currency, warnings)
            balances.append(
                {
                    "broker": BROKER_CODE,
                    "account_id": account_id,
                    "year": report.period_end.year,
                    "date": report.period_end.isoformat(),
                    "currency": currency,
                    "ending_cash": _money_text(ending_cash),
                    "ending_cash_kzt": _amount_kzt(ending_cash, rate),
                    "source_report": row.get("source_report"),
                }
            )
    return balances


def _populate_raw_totals(
    dataset: CanonicalDataset,
    reports: Sequence[ParsedHalykReport],
    trades: Sequence[Mapping[str, Any]],
) -> None:
    total_amount = Decimal("0")
    total_commission = Decimal("0")
    turnover_metric = ReconciliationMetric.TRADE_GROSS_AMOUNT_BY_INSTRUMENT.value
    for trade in trades:
        amount = abs(_decimal(trade.get("amount")))
        commission = abs(_decimal(trade.get("commission")))
        total_amount += amount
        total_commission += commission
        key = _dimension_key(
            metric=turnover_metric,
            year=_year(trade.get("date_time")),
            currency=_currency(trade.get("currency")),
            instrument_key=_text(trade.get("isin") or trade.get("symbol")),
        )
        dataset.raw_totals.totals_by_metric_currency[key] = (
            dataset.raw_totals.totals_by_metric_currency.get(key, Decimal("0")) + amount
        )
    dataset.raw_totals.scalar_totals[ReconciliationMetric.TOTAL_TRADES_GROSS_AMOUNT.value] = total_amount
    dataset.raw_totals.scalar_totals[ReconciliationMetric.TOTAL_COMMISSIONS.value] = total_commission
    dataset.raw_totals.scalar_totals[ReconciliationMetric.TOTAL_DIVIDENDS_GROSS.value] = sum(
        (_decimal(row.get("gross_amount")) for row in dataset.tables.get("Dividends", [])), Decimal("0")
    )
    dataset.raw_totals.scalar_totals[ReconciliationMetric.TOTAL_DIVIDENDS_TAX.value] = Decimal("0")
    dataset.raw_totals.scalar_totals[ReconciliationMetric.TOTAL_DIVIDENDS_NET.value] = dataset.raw_totals.scalar_totals[
        ReconciliationMetric.TOTAL_DIVIDENDS_GROSS.value
    ]
    dataset.raw_totals.scalar_totals[ReconciliationMetric.TOTAL_COUPONS.value] = sum(
        (_decimal(row.get("gross_amount")) for row in dataset.tables.get("Coupons", [])), Decimal("0")
    )
    dataset.raw_totals.scalar_totals[ReconciliationMetric.TOTAL_INTEREST.value] = sum(
        (_decimal(row.get("gross_amount")) for row in dataset.tables.get("Interest", [])), Decimal("0")
    )

    for row in dataset.tables.get("CashBalances", []):
        key = _dimension_key(year=int(row["year"]), currency=_currency(row.get("currency")))
        dataset.raw_totals.cash_by_currency[key] = (
            dataset.raw_totals.cash_by_currency.get(key, Decimal("0")) + _decimal(row.get("ending_cash"))
        )
    for report in reports:
        if report.period_end is None:
            continue
        for row in report.positions:
            isin = _text(row.get("isin"))
            if not isin:
                continue
            key = _dimension_key(year=report.period_end.year, instrument_key=isin)
            dataset.raw_totals.positions_by_key[key] = (
                dataset.raw_totals.positions_by_key.get(key, Decimal("0")) + _decimal(row.get("quantity"))
            )


def _movement_kind(operation: str) -> str | None:
    normalized = operation.casefold()
    if "при покупк" in normalized:
        return "purchase"
    if "при продаж" in normalized:
        return "sale"
    if "погаш" in normalized:
        return "redemption"
    return None


def _is_self_security_transfer_in(operation: str, isin: str | None, converted_isin: str | None) -> bool:
    """Identify Halyk's custody credit order for the same security, not a conversion."""

    normalized = operation.casefold()
    return bool(
        isin
        and converted_isin
        and isin == converted_isin
        and "приказ на зачислен" in normalized
    )


def _movement_exchange(values: Sequence[Any], headers: Mapping[str, int]) -> str | None:
    market = _text(_column(values, headers, "Рынок ЦБ"))
    if market:
        return _normalize_exchange(market)
    transfer_from = _text(_column(values, headers, "Перевод откуда"))
    transfer_to = _text(_column(values, headers, "Перевод куда"))
    counterpart = transfer_from if transfer_from not in {None, "-"} else transfer_to
    return "OTC" if counterpart not in {None, "-"} else None


def _commission_in_trade_currency(
    commission_kzt: Decimal,
    year: int | None,
    currency: str,
    fx_provider: AnnualFxRateProvider,
    warnings: list[str],
) -> Decimal:
    if commission_kzt == 0 or currency == BASE_CURRENCY:
        return commission_kzt
    rate = _annual_rate(fx_provider, year, currency, warnings)
    if rate is None or rate == 0:
        return Decimal("0")
    return commission_kzt / rate


def _unprocessed(
    row: Mapping[str, Any], reason: str, details: str
) -> dict[str, Any]:
    source_report = _text(row.get("source_report"))
    return {
        "severity": "warning",
        "reason": reason,
        "details": details,
        "source_sheet": "Trades",
        "source_report": source_report,
        "trade_id": f"{Path(source_report).name}:{row.get('source_row')}" if source_report else None,
        "date_time": row.get("date_time") or row.get("date"),
        "symbol": row.get("symbol"),
        "isin": row.get("isin"),
        "asset_type": _asset_type(row.get("security_type")),
        "currency": row.get("currency"),
        "quantity": row.get("quantity"),
        "price": row.get("price"),
        "amount": row.get("amount"),
        "commission": row.get("commission") or row.get("commission_kzt"),
    }


def _column(values: Sequence[Any], headers: Mapping[str, int], name: str) -> Any:
    index = headers.get(name)
    return values[index] if index is not None and index < len(values) else None


def _is_data_number(value: Any) -> bool:
    if isinstance(value, (int, float)):
        return True
    try:
        Decimal(str(value))
        return True
    except (InvalidOperation, ValueError, TypeError):
        return False


def _parse_period(value: Any) -> tuple[date | None, date | None]:
    matches = _DATE_PATTERN.findall(_clean_text(value))
    if not matches:
        report_date = _parse_date(value)
        return report_date, report_date
    dates = [_parse_date(item) for item in matches]
    valid = [item for item in dates if item is not None]
    if not valid:
        return None, None
    return valid[0], valid[-1]


def _parse_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _clean_text(value)
    for fmt in ("%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text[:10], fmt).date()
        except ValueError:
            pass
    return None


def _date_text(value: Any) -> str | None:
    parsed = _parse_date(value)
    return parsed.isoformat() if parsed else None


def _date_part(value: Any) -> str | None:
    text = _clean_text(value)
    return text[:10] if len(text) >= 10 else None


def _decimal(value: Any) -> Decimal:
    if value in (None, ""):
        return Decimal("0")
    text = str(value).replace("\u00a0", "").replace(" ", "").replace(",", ".").strip()
    try:
        return Decimal(text or "0")
    except InvalidOperation:
        return Decimal("0")


def _clean_text(value: Any) -> str:
    return " ".join(str(value or "").replace("\u00a0", " ").split()).strip()


def _text(value: Any) -> str | None:
    text = _clean_text(value)
    return text or None


def _currency(value: Any) -> str | None:
    text = _text(value)
    return text.upper() if text else None


def _symbol_from_name(value: Any, isin: str | None) -> str:
    symbol = _clean_text(value).split(",", 1)[0].strip()
    return symbol or isin or "UNKNOWN"


def _asset_type(value: Any) -> str:
    normalized = _clean_text(value).casefold()
    return "Bonds" if "облигац" in normalized else "Stocks"


def _normalize_exchange(value: Any) -> str | None:
    exchange = _text(value)
    if not exchange:
        return None
    upper = exchange.upper()
    if "KASE" in upper:
        return "KASE"
    if "AIX" in upper:
        return "AIX"
    if upper == "NON EXCHANGE":
        return "OTC"
    return exchange


def _country(raw_country: Any, isin: str | None) -> str | None:
    country = _text(raw_country)
    if country:
        return country.upper()
    return isin[:2].upper() if isin and len(isin) >= 2 else None


def _same_account_id(left: Any, right: Any) -> bool:
    left_text = _clean_text(left)
    right_text = _clean_text(right)
    if left_text.isdigit() and right_text.isdigit():
        return left_text.lstrip("0") == right_text.lstrip("0")
    return left_text.casefold() == right_text.casefold()


def _max_report_year(reports: Sequence[ParsedHalykReport]) -> int | None:
    return max((report.period_end.year for report in reports if report.period_end), default=None)


def _year(value: Any) -> int | None:
    parsed = _parse_date(value)
    return parsed.year if parsed else None


def _annual_rate(
    fx_provider: AnnualFxRateProvider,
    year: int | None,
    currency: str | None,
    warnings: list[str],
) -> Decimal | None:
    if year is None or not currency:
        return None
    rate = fx_provider.rate(year, currency)
    if rate is None:
        warning = f"Missing annual NBK FX rate for {currency}/{year}; KZT fields left empty."
        if warning not in warnings:
            warnings.append(warning)
    return rate


def _money_text(value: Decimal) -> str:
    return str(value.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def _amount_kzt(amount: Decimal, rate: Decimal | None) -> str | None:
    return str(amount * rate) if rate is not None else None


def _earlier_date(left: date | None, right: date | None) -> date | None:
    if left is None:
        return right
    if right is None:
        return left
    return min(left, right)


def _later_date(left: date | None, right: date | None) -> date | None:
    if left is None:
        return right
    if right is None:
        return left
    return max(left, right)


def _dimension_key(
    *,
    metric: str | None = None,
    year: int | None = None,
    currency: str | None = None,
    instrument_key: str | None = None,
) -> str:
    return "|".join("" if value is None else str(value) for value in (metric, year, currency, instrument_key))
